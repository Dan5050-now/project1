"""Work the FTE calculation through, scenario by scenario, using the REAL engine.

Every figure in docs/PRAP_FTE_계산설명서.pdf comes from here, and here runs
tools/prap_io.py - the same reference implementation the browser is held to on every
test run. Nothing in the document is hand-computed, so it cannot drift from what the
application would actually produce.

One small project is built from the delivered template's own standards and role
factors, then varied one thing at a time:

    base    세 사람, 세 역할
    A       같은 역할에 인원이 늘어날 때
    B       개인 가중치를 낮출 때
    C       기간 가중치를 조정할 때
    D       달의 일부만 참여할 때
    E       배정 단위로 월 FTE를 수동 지정할 때
    F       과제 단위로 월 FTE를 수동 지정할 때
    G       역할에 아무도 없을 때 (흡수)
    H       개인 capacity 를 낮출 때

    python tools/fte_examples.py            표로 출력
    python tools/fte_examples.py --json     문서 생성기가 읽는 형태로

Output: output/deck/fte_examples.json
"""

import copy
import json
import pathlib
import subprocess
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
TPL = ROOT / "templates" / "PRAP_SourceData_Dummy_10x10_v1.10.xlsx"
OUT = ROOT / "output" / "deck"
TMP = ROOT / "output" / "deck" / "_fte"

PID = "PRJ-X01"
TYPE, PHASE, SCOPE = "NewDrug CT", "Phase 3", "fully in-housed"
PERIOD = "Conduct (final)"
START, END = datetime(2026, 1, 1), datetime(2026, 12, 31)
MONTH = "2026-06"                       # the month every table below reports
MONTH_LABEL = "Jun 2026"                # how the engine's CLI writes it
ROLES = ["Lead data manager", "Clinical Data Associator", "Data Analyst"]
# Every person who appears in ANY scenario, so the parser can name them all.
ALL_PEOPLE = [("PSN-X1", "김 O O", "Lead data manager"),
              ("PSN-X2", "이 O O", "Clinical Data Associator"),
              ("PSN-X3", "박 O O", "Data Analyst"),
              ("PSN-X4", "최 O O", "Lead data manager")]

PEOPLE = [("PSN-X1", "김 O O", "Lead data manager"),
          ("PSN-X2", "이 O O", "Clinical Data Associator"),
          ("PSN-X3", "박 O O", "Data Analyst")]


def base_rows():
    """The starting plan: one project, one period, three people in three roles."""
    proj = dict(project_id=PID, project_name="예시 과제", project_type=TYPE,
                clinical_phase=PHASE, work_scope_type=SCOPE,
                start_date=START, end_date=END, status="Active",
                estimation_type="automatic")
    per = [dict(project_id=PID, period_name=PERIOD, period_seq=1,
                period_start=START, period_end=END, weight=1.00)]
    ppl = [dict(person_id=p, person_name=n, department="Data Management",
                primary_role=r, capacity_fte=1.00) for p, n, r in PEOPLE]
    asg = [dict(assignment_id=f"ASG-X{i}", person_id=p, project_id=PID, role_name=r,
                person_weight=1.00, estimation_type="automatic")
           for i, (p, n, r) in enumerate(PEOPLE, 1)]
    return proj, per, ppl, asg


def build(path, proj, per, ppl, asg, ppw=None, est=None):
    """Write a workbook that carries these rows plus the template's own reference data."""
    src = load_workbook(TPL)
    wb = Workbook()
    wb.remove(wb.active)
    order = ["Project", "Milestone", "ProjectPeriod", "PeriodFTEStandard", "RoleFactor",
             "Person", "Assignment", "PersonPeriodWeight", "MonthlyEstimate", "Lists",
             "Config"]
    given = {"Project": [proj], "ProjectPeriod": per, "Person": ppl, "Assignment": asg,
             "PersonPeriodWeight": ppw or [], "MonthlyEstimate": est or [],
             "Milestone": []}
    for name in order:
        s, d = src[name], wb.create_sheet(name)
        hdr = [c.value for c in s[1]]
        d.append(hdr)
        if name in given:
            for row in given[name]:
                d.append([row.get(h) for h in hdr])
        else:                                   # reference sheets, copied verbatim
            for r in s.iter_rows(min_row=2, values_only=True):
                if any(v is not None for v in r):
                    d.append(list(r))
    wb.save(path)


def run(path):
    """prap_io's own per-cell figures for MONTH - the engine, not a re-derivation."""
    jf = path.with_suffix(".json")
    subprocess.run([sys.executable, str(ROOT / "tools" / "prap_io.py"), "to-json",
                    str(path), "-o", str(jf)], check=True, capture_output=True)
    out = subprocess.run([sys.executable, str(ROOT / "tools" / "prap_io.py"), "calculate",
                          str(jf), "--by", "cell"], capture_output=True, text=True).stdout
    # The CLI prints a heading per cell - project, person, role - then one indented line
    # per month reading "Jun 2026   1.26". So track the heading and read the month under
    # it; matching on the heading alone would attribute every month to the first cell.
    label = MONTH_LABEL
    rows, total, cur = [], 0.0, None
    for line in out.splitlines():
        if not line.strip():
            continue
        if not line.startswith("      "):                   # a cell heading
            cur = line.strip()
            continue
        if cur and line.strip().startswith(label):
            try:
                val = float(line.strip().split()[-1])
            except ValueError:
                continue
            pid = next((p for p, n, r in ALL_PEOPLE if p in cur), None)
            rows.append({
                "person": next((n for p, n, r in ALL_PEOPLE if p == pid), None),
                "role": next((r for r in ROLES if r in cur), None),
                "fte": round(val, 2)})
            total += val
    return rows, round(total, 2)


SCENARIOS = []


def scenario(key, title, question, note, mutate):
    SCENARIOS.append((key, title, question, note, mutate))


scenario("base", "기준 상황", "세 사람이 각각 다른 역할을 맡는다",
         "과제가 그 달에 필요로 하는 양을 세 사람이 역할 비중에 따라 나눈다.",
         lambda p, pe, pl, a: None)

scenario("A", "같은 역할에 인원이 늘어날 때",
         "Lead data manager 에 한 사람을 더 넣으면?",
         "역할에 매겨진 비중은 그대로이고, 그 역할을 맡은 사람들이 나눠 갖는다. "
         "과제가 필요로 하는 양은 변하지 않는다 — 사람을 더 넣는다고 일이 늘지 않는다.",
         lambda p, pe, pl, a: (
             pl.append(dict(person_id="PSN-X4", person_name="최 O O",
                            department="Data Management",
                            primary_role="Lead data manager", capacity_fte=1.00)),
             a.append(dict(assignment_id="ASG-X4", person_id="PSN-X4", project_id=PID,
                           role_name="Lead data manager", person_weight=1.00,
                           estimation_type="automatic"))))

scenario("B", "개인 가중치를 낮출 때",
         "김 O O 의 참여 비중을 1.00 에서 0.50 으로 낮추면?",
         "본인 몫은 줄지만 과제가 필요로 하는 양은 그대로이므로, 줄어든 만큼이 "
         "나머지 사람에게 넘어간다. 인력이 모자란 상황이 과제가 아니라 사람 쪽에 나타난다.",
         lambda p, pe, pl, a: a[0].update(person_weight=0.50))

scenario("C", "기간 가중치를 조정할 때",
         "이 과제가 같은 종류의 보통 과제보다 20% 무겁다면?",
         "기간 가중치는 과제가 필요로 하는 양 자체를 키우거나 줄인다. "
         "나누는 비율은 그대로이므로 모든 사람의 몫이 같은 비율로 늘어난다.",
         lambda p, pe, pl, a: pe[0].update(weight=1.20))

scenario("D", "달의 일부만 참여할 때",
         "박 O O 가 6월 16일에 합류하면?",
         "그 달에 실제로 걸쳐 있는 날수만큼만 몫을 주장한다. "
         "남은 몫은 그 달 내내 일한 사람들에게 돌아간다.",
         lambda p, pe, pl, a: a[2].update(assign_start_date=datetime(2026, 6, 16)))

scenario("E", "배정 단위로 월 FTE 를 수동 지정할 때",
         "이 O O 의 6월 값을 2.00 으로 직접 적으면?",
         "그 사람의 몫은 계산 대신 적은 값이 된다. 다른 사람의 몫은 계산한 그대로이므로, "
         "과제의 그 달 합계는 표준과 달라질 수 있다 — 그 차이는 앱이 따로 알려준다.",
         None)

scenario("F", "과제 단위로 월 FTE 를 수동 지정할 때",
         "과제의 6월 전체를 5.00 으로 적으면?",
         "적은 값이 그 달 전체가 되고, 그 달에 배정된 사람들이 원래 비율 그대로 "
         "다시 나눠 갖는다. 과제 합계와 사람 합계는 항상 일치한다.",
         None)

scenario("G", "역할에 아무도 없을 때",
         "Clinical Data Associator 를 배정하지 않으면?",
         "그 역할의 일이 사라지지는 않는다. 기준표에 '이 역할이 비면 누가 대신 받는지'가 "
         "적혀 있고, 여기서는 Lead data manager 가 받는다. 그래서 그 사람의 몫이 커진다.",
         lambda p, pe, pl, a: (pl.pop(1), a.pop(1)))

scenario("H", "개인 capacity 를 낮출 때",
         "김 O O 의 capacity 를 0.50 으로 낮추면?",
         "계산은 달라지지 않는다. capacity 는 '이 사람이 쓸 수 있는 시간이 얼마인가'를 "
         "적어두는 값이고, 몫을 정하는 데 쓰이지 않는다. "
         "과부하 판정 기준도 capacity 가 아니라 정해진 절대값과 비교한다.",
         lambda p, pe, pl, a: pl[0].update(capacity_fte=0.50))


def months_of_project():
    out = []
    y, m = START.year, START.month
    while (y, m) <= (END.year, END.month):
        out.append(f"{y}-{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def main(as_json=False):
    TMP.mkdir(parents=True, exist_ok=True)
    results = {}
    for key, title, question, note, mutate in SCENARIOS:
        proj, per, ppl, asg = base_rows()
        ppw, est = [], []
        if mutate:
            mutate(proj, per, ppl, asg)
        if key == "E":
            asg[1]["estimation_type"] = "manual"
            for mo in months_of_project():
                est.append(dict(scope="assignment", ref_id="ASG-X2", month=mo,
                                fte=2.00 if mo == MONTH else 1.00))
        if key == "F":
            proj["estimation_type"] = "manual"
            for mo in months_of_project():
                est.append(dict(scope="project", ref_id=PID, month=mo,
                                fte=5.00 if mo == MONTH else 3.47))
        path = TMP / f"{key}.xlsx"
        build(path, proj, per, ppl, asg, ppw, est)
        rows, total = run(path)
        results[key] = {"title": title, "question": question, "note": note,
                        "rows": rows, "total": total}
        if not as_json:
            print(f"\n=== [{key}] {title}")
            print(f"    {question}")
            for r in rows:
                print(f"      {r['person'] or '?':10} {r['role'] or '':32} {r['fte']:>6.2f}")
            print(f"      {'과제 합계':10} {'':32} {total:>6.2f}")
    if as_json:
        OUT.mkdir(parents=True, exist_ok=True)
        (OUT / "fte_examples.json").write_text(
            json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"written: output/deck/fte_examples.json  ({len(results)} scenarios)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main("--json" in sys.argv))
