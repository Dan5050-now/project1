"""A scorecard for a resource plan, computed in code so an AI only has to explain it.

The calculated-FTE export answers "what is each project given, and what does it need".
Somebody analysing it - a person or a language model - then wants a verdict: how healthy
is this plan, where is it short, who is overloaded, what is coming, and what could be
moved. A language model asked to work that out from four thousand rows will add things
up wrong and give a different score every time it is asked. So the arithmetic is done
HERE, deterministically, and the model is handed the result to interpret:

    python tools/prap_analyze.py score   <file>          [--as-of YYYY-MM] [--out DIR]
    python tools/prap_analyze.py compare <before> <after> [--as-of YYYY-MM] [--out DIR]

<file> is either the CALCULATED export ("Export calculated FTE", app 1.57 or later) or a
SOURCE plan (.xlsx or .prap.json). A source plan is calculated with tools/prap_io.py,
the reference implementation that tools/test_interop.py holds to the application - so a
what-if can be drafted as a source file and scored without a browser, by the same
formula as the export it is compared with.

Both commands write <name>.json (for a program or an AI) and <name>.md (for a person),
and print the markdown. --pseudonymise drops person names and replaces departments with
codes, writing the key to a separate local file that is NOT for sharing.

What the score is, and what it is not: a fixed formula, below in SCORING, with every
component's inputs written out beside its points. It is a way of comparing a plan with
itself over time and with a scenario - not an absolute grade, and the weights are a
starting point to agree with the people who use it. Change them here, once, and bump
SCORING["version"] so two scorecards made under different weights are never compared.

Candidate moves are INDICATIVE. The engine divides a project's demand among whoever is
on it (REQ-CAL-19), so adding a person changes everybody's share; spare capacity says
where to look, not what the figures will be. Confirm any move by editing the source and
running `compare`, which is what docs/PRAP_AI_Analysis_Guide.md tells an AI to do.
"""

import argparse
import copy
import json
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import prap_io                                                          # noqa: E402

# ---------------------------------------------------------------------------- scoring
# Every component is "how much of X is in trouble", over a denominator stated beside it,
# and scores   weight x 0.5 ^ (ratio / half_at)   - full marks at 0, half at `half_at`,
# a quarter at twice that. A curve that never reaches zero, deliberately: a straight
# line to a floor goes flat once a plan is past it, and then a scenario that takes a
# person from 21 months over the ceiling to 8 scores exactly the same as doing nothing.
# A score that cannot see an improvement is no use for comparing two plans.
SCORING = {
    "version": "1",
    "components": {
        "demand_coverage": {
            "weight": 35, "half_at": 0.05,
            "measures": "FTE-months short of the standard or unallocated, over the "
                        "standard demand. A project with too few people is not short "
                        "here - its people go over the ceiling instead, which "
                        "over_allocation scores and the project list shows as hidden",
        },
        "over_allocation": {
            "weight": 25, "half_at": 0.10,
            "measures": "FTE-months above the over-allocation ceiling, over the FTE-months "
                        "staffed - how much of the work has nobody with room to do it",
        },
        "concentration": {
            "weight": 15, "half_at": 0.15,
            "measures": "the mean of two ratios: project-months of at least KEY_DEMAND FTE "
                        "that one person carries more than KEY_SHARE of, and person-months "
                        "spread over FRAGMENTED or more projects",
        },
        "plan_reliability": {
            "weight": 15, "half_at": 0.05,
            "measures": "FTE stated by hand that departs from what the assumptions give "
                        "(sum of |fte - automatic_fte|), over the automatic total",
        },
        "under_use": {
            "weight": 10, "half_at": 0.10,
            "measures": "person-months inside an under-allocation run, over the "
                        "person-months anybody is staffed",
        },
    },
    "bands": [(85, "healthy"), (70, "watch"), (50, "at risk"), (0, "critical")],
}
KEY_DEMAND = 1.0       # a project-month this big should not rest on one person...
KEY_SHARE = 0.60       # ...carrying more than this share of it
FRAGMENTED = 5         # a person on this many projects in one month is spread thin
LOOKAHEAD = 6          # months the forward look covers
MIN_HEADROOM = 0.10    # spare capacity worth proposing
TOP = 10               # rows per ranked list


def cents(v):
    return prap_io.to_cents(v or 0)


def fc(c):
    return round(c / prap_io.CENTS, 2)


def iso_to_k(s):
    y, m = str(s)[:7].split("-")
    return int(y) * 12 + int(m) - 1


def k_to_iso(k):
    return f"{k // 12}-{k % 12 + 1:02d}"


# ------------------------------------------------------------------------- the data
class Plan:
    """One plan, as five lists of plain rows - whichever kind of file it came from.

    pm     project-months: demand, staffed, gap, dir  (dir short/over/unallocated/None)
    psm    person-months:  fte, capacity
    lines  assignment-months: project, person, role, fte, auto
    """

    def __init__(self):
        self.pm, self.psm, self.lines = [], [], []
        self.projects, self.people = {}, {}
        self.months = []
        self.OVER, self.UNDER, self.MINM = 1.5, 0.6, 3
        self.source, self.kind, self.filters = "", "", "none"

    def restrict(self, months):
        keep = set(months)
        self.months = sorted(keep)
        self.pm = [r for r in self.pm if r["k"] in keep]
        self.psm = [r for r in self.psm if r["k"] in keep]
        self.lines = [r for r in self.lines if r["k"] in keep]


def _sheet(wb, name):
    it = wb[name].iter_rows(values_only=True)
    hdr = next(it)
    return [dict(zip(hdr, r)) for r in it if any(v is not None for v in r)]


def read_export(path, wb):
    P = Plan()
    P.kind, P.source = "calculated export", Path(path).name
    pm = _sheet(wb, "ProjectMonth")
    if pm and "demand_fte" not in pm[0]:
        raise SystemExit(
            f"error: {Path(path).name} was exported before the application carried each "
            "project's demand and gap (app 1.57, R-55). Export it again from a current "
            "application, or pass the SOURCE plan instead.")
    for r in pm:
        k = iso_to_k(r["month_iso"])
        P.projects.setdefault(r["project_id"], {
            "project_name": r["project_name"], "project_type": r["project_type"],
            "clinical_phase": r["clinical_phase"], "status": r["status"]})
        P.pm.append({"k": k, "pid": r["project_id"], "demand": r["demand_fte"] or 0,
                     "staffed": r["staffed_fte"] or 0, "gap": r["gap_fte"] or 0,
                     "dir": r["gap_dir"] or None, "period": r["period_name"]})
    for r in _sheet(wb, "PersonMonth"):
        P.people.setdefault(r["person_id"], {
            "person_name": r["person_name"], "department": r["department"],
            "primary_role": r["primary_role"], "capacity_fte": r["capacity_fte"]})
        P.psm.append({"k": iso_to_k(r["month_iso"]), "sid": r["person_id"],
                      "fte": r["fte"] or 0})
    for r in _sheet(wb, "Detail"):
        P.people.setdefault(r["person_id"], {
            "person_name": r["person_name"], "department": r["department"],
            "primary_role": None, "capacity_fte": None})
        P.lines.append({"k": iso_to_k(r["month_iso"]), "pid": r["project_id"],
                        "sid": r["person_id"], "role": r["role_name"],
                        "fte": r["fte"] or 0, "auto": r["automatic_fte"] or 0})
    cfg = {r["parameter"]: r["value"] for r in _sheet(wb, "Assumptions")}
    P.OVER = float(cfg.get("over_allocation_fte") or 1.5)
    P.UNDER = float(cfg.get("under_allocation_fte") or 0.6)
    P.MINM = int(float(cfg.get("under_allocation_min_months") or 3))
    rows = list(wb["00_ReadMe"].iter_rows(values_only=True))
    for r in rows:
        if r and r[0] == "Filters" and len(r) > 1:
            P.filters = str(r[1])
    P.months = sorted({r["k"] for r in P.pm} | {r["k"] for r in P.psm})
    return P


def read_source(path):
    P = Plan()
    P.kind, P.source = "source plan", Path(path).name
    sheets = prap_io.read_any(path)
    pristine = copy.deepcopy(sheets)          # Model() adds derived rows to what it is given
    M = prap_io.Model(sheets)
    C = prap_io.calculate(M)
    # What the assumptions alone would give, line by line: the same plan with nothing
    # stated by hand. The export carries this as automatic_fte; a source file has to be
    # calculated twice to get it.
    M0 = prap_io.Model(pristine)
    M0.manual = {}
    C0 = prap_io.calculate(M0)
    auto = defaultdict(float)
    for L in C0["lines"]:
        auto[(L.get("assignment_id"), L["month"])] += L["fte"]

    for pid, p in M.projects.items():
        P.projects[pid] = {k: p.get(k) for k in
                           ("project_name", "project_type", "clinical_phase", "status")}
    for sid, p in M.people.items():
        P.people[sid] = {k: p.get(k) for k in
                         ("person_name", "department", "primary_role", "capacity_fte")}
    P.OVER, P.UNDER, P.MINM = M.OVER, M.UNDER, M.MINM

    for L in C["lines"]:
        P.lines.append({"k": L["month"], "pid": L["project_id"], "sid": L["person_id"],
                        "role": L.get("role_name"), "fte": L["fte"],
                        "auto": auto.get((L.get("assignment_id"), L["month"]), 0.0)})
    for (sid, k), v in C["pers_month"].items():
        if cents(v) > 0:
            P.psm.append({"k": k, "sid": sid, "fte": fc(cents(v))})
    gap = C["proj_gap"]
    for (pid, k), v in C["proj_month"].items():
        g = gap.get((pid, k))
        staffed = fc(cents(v))
        demand = fc(cents(g["demand"])) if g else staffed
        P.pm.append({"k": k, "pid": pid, "demand": demand, "staffed": staffed,
                     "gap": fc(cents(staffed) - cents(demand)),
                     "dir": g["dir"] if g else None, "period": None})
    for (pid, k), d in C["proj_unallocated"].items():
        P.pm.append({"k": k, "pid": pid, "demand": fc(cents(d)), "staffed": 0.0,
                     "gap": -fc(cents(d)), "dir": "unallocated", "period": None})
    P.months = sorted({r["k"] for r in P.pm} | {r["k"] for r in P.psm})
    # As the export does it (R-55): a month somebody is employed and on nothing is a
    # row with fte 0, so spare capacity is visible and a leaver is never offered as it.
    have = {(r["sid"], r["k"]) for r in P.psm}
    for sid, p in M.people.items():
        for k in range(P.months[0], P.months[-1] + 1) if P.months else ():
            if (sid, k) not in have and employed(p, k):
                P.psm.append({"k": k, "sid": sid, "fte": 0.0})
    return P


def employed(p, k):
    y, m = divmod(k, 12)
    first = date(y, m + 1, 1)
    last = date(y + (m + 1) // 12, (m + 1) % 12 + 1, 1)          # first of next month
    s, e = p.get("employment_start"), p.get("employment_end")
    s = s.date() if hasattr(s, "date") else s
    e = e.date() if hasattr(e, "date") else e
    return not (isinstance(s, date) and s >= last) and not (isinstance(e, date) and e < first)


def load(path):
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        wb = load_workbook(p, read_only=True, data_only=True)
        if "00_ReadMe" in wb.sheetnames and "ProjectMonth" in wb.sheetnames:
            return read_export(p, wb)
        wb.close()
    return read_source(p)


def horizon(P, frm=None, to=None):
    lo = iso_to_k(frm) if frm else None
    hi = iso_to_k(to) if to else None
    if lo is not None or hi is not None:
        P.restrict([k for k in P.months
                    if (lo is None or k >= lo) and (hi is None or k <= hi)])


# --------------------------------------------------------------------------- analysis
def runs_of(keys):
    """Consecutive month keys, as [(first, last, n)]."""
    out, cur = [], None
    for k in sorted(keys):
        if cur and k == cur[1] + 1:
            cur = (cur[0], k, cur[2] + 1)
        else:
            if cur:
                out.append(cur)
            cur = (k, k, 1)
    if cur:
        out.append(cur)
    return out


def span(r):
    a, b, n = r
    return k_to_iso(a) if n == 1 else f"{k_to_iso(a)} to {k_to_iso(b)}"


def band(score):
    return next(name for floor, name in SCORING["bands"] if score >= floor)


def analyse(P, as_of, top=TOP):
    pname = lambda pid: (P.projects.get(pid) or {}).get("project_name") or pid
    sname = lambda sid: (P.people.get(sid) or {}).get("person_name") or sid

    # ---- figures every component draws on ------------------------------------------
    demand_c = sum(cents(r["demand"]) for r in P.pm)
    short_c = sum(-cents(r["gap"]) for r in P.pm if r["dir"] == "short")
    unall_c = sum(cents(r["demand"]) for r in P.pm if r["dir"] == "unallocated")
    over_std_c = sum(cents(r["gap"]) for r in P.pm if r["dir"] == "over")
    staffed_pm = [r for r in P.psm if cents(r["fte"]) > 0]
    over_pm = [r for r in staffed_pm if r["fte"] > P.OVER + 1e-9]

    fte_of = {(r["sid"], r["k"]): r["fte"] for r in P.psm}
    under_runs, under_months = [], 0
    for sid in sorted({r["sid"] for r in P.psm}):
        ks = [k for k in P.months if 0 < fte_of.get((sid, k), 0) < P.UNDER - 1e-9]
        for run in runs_of(ks):
            if run[2] >= P.MINM:
                under_runs.append((sid, run))
                under_months += run[2]

    by_pm = defaultdict(list)                    # (pid, k) -> lines
    for L in P.lines:
        by_pm[(L["pid"], L["k"])].append(L)
    key_person = []
    big = 0
    for (pid, k), ls in by_pm.items():
        tot = sum(L["fte"] for L in ls)
        if tot < KEY_DEMAND - 1e-9:
            continue
        big += 1
        per = defaultdict(float)
        for L in ls:
            per[L["sid"]] += L["fte"]
        sid, most = max(per.items(), key=lambda kv: kv[1])
        if most / tot > KEY_SHARE + 1e-9:
            key_person.append((pid, k, sid, most / tot))
    nproj = defaultdict(set)
    for L in P.lines:
        if L["fte"] > 0:
            nproj[(L["sid"], L["k"])].add(L["pid"])
    fragmented = [key for key, s in nproj.items() if len(s) >= FRAGMENTED]

    auto_c = sum(cents(L["auto"]) for L in P.lines)
    depart_c = sum(abs(cents(L["fte"]) - cents(L["auto"])) for L in P.lines)

    # ---- the components ---------------------------------------------------------------
    def comp(name, ratio, inputs):
        spec = SCORING["components"][name]
        frac = 0.5 ** (ratio / spec["half_at"]) if ratio is not None else 1.0
        return {"points": round(spec["weight"] * frac, 1), "of": spec["weight"],
                "ratio": None if ratio is None else round(ratio, 4),
                "half_at": spec["half_at"], "measures": spec["measures"],
                "inputs": inputs}

    n_staffed = len(staffed_pm)
    kp_ratio = len(key_person) / big if big else 0.0
    fr_ratio = len(fragmented) / n_staffed if n_staffed else 0.0
    conc_ratio = (kp_ratio + fr_ratio) / 2
    staffed_c = sum(cents(r["fte"]) for r in staffed_pm)
    excess_c = sum(cents(r["fte"]) - cents(P.OVER) for r in over_pm)
    comps = {
        "demand_coverage": comp(
            "demand_coverage", (short_c + unall_c) / demand_c if demand_c else None,
            {"standard_demand_fte_months": fc(demand_c),
             "short_fte_months": fc(short_c), "unallocated_fte_months": fc(unall_c)}),
        "over_allocation": comp(
            "over_allocation", excess_c / staffed_c if staffed_c else None,
            {"excess_fte_months": fc(excess_c), "staffed_fte_months": fc(staffed_c),
             "over_person_months": len(over_pm), "staffed_person_months": n_staffed,
             "ceiling_fte": P.OVER}),
        "concentration": comp(
            "concentration", conc_ratio,
            {"key_person_project_months": len(key_person),
             f"project_months_of_{KEY_DEMAND}_fte_or_more": big,
             "key_person_ratio": round(kp_ratio, 4),
             f"person_months_on_{FRAGMENTED}_or_more_projects": len(fragmented),
             "fragmented_ratio": round(fr_ratio, 4)}),
        "plan_reliability": comp(
            "plan_reliability", depart_c / auto_c if auto_c else None,
            {"departure_fte_months": fc(depart_c), "automatic_fte_months": fc(auto_c),
             "over_standard_fte_months": fc(over_std_c)}),
        "under_use": comp(
            "under_use", under_months / n_staffed if n_staffed else None,
            {"person_months_in_under_runs": under_months, "under_runs": len(under_runs),
             "floor_fte": P.UNDER, "min_months": P.MINM}),
    }
    score = round(sum(c["points"] for c in comps.values()), 1)

    # ---- the shortfall the engine does not show as a gap --------------------------
    # A project's month IS its demand, shared among whoever is on it (REQ-CAL-19). So a
    # project with too few people is never short on paper: its people go over the
    # ceiling instead. That overload is the project's real shortfall, and leaving it on
    # the people would let a plan score full marks for coverage while most of its staff
    # are carrying twice what they can. So each over-ceiling person-month's EXCESS is
    # attributed back to the projects making it up, in proportion to what each
    # contributes - "hidden" because no gap column will ever show it.
    contrib = defaultdict(lambda: defaultdict(float))     # (sid, k) -> pid -> fte
    for L in P.lines:
        contrib[(L["sid"], L["k"])][L["pid"]] += L["fte"]
    hidden = defaultdict(float)                           # (pid, k) -> FTE nobody has
    for r in over_pm:
        excess = r["fte"] - P.OVER
        for pid, v in contrib[(r["sid"], r["k"])].items():
            if v > 0:
                hidden[(pid, r["k"])] += excess * v / r["fte"]

    # ---- projects, ranked by what they are missing ---------------------------------
    proj = defaultdict(lambda: {"short": 0, "unallocated": 0, "over": 0, "hidden": 0.0,
                                "months": defaultdict(list), "demand": 0})
    for (pid, k), v in hidden.items():
        if round(v, 2) > 0:
            proj[pid]["hidden"] += v
            proj[pid]["months"]["hidden"].append(k)
    for r in P.pm:
        e = proj[r["pid"]]
        e["demand"] += cents(r["demand"])
        if r["dir"] == "short":
            e["short"] += -cents(r["gap"])
        elif r["dir"] == "unallocated":
            e["unallocated"] += cents(r["demand"])
        elif r["dir"] == "over":
            e["over"] += cents(r["gap"])
        if r["dir"]:
            e["months"][r["dir"]].append(r["k"])
    projects = []
    for pid, e in proj.items():
        if not (e["short"] or e["unallocated"] or e["over"] or round(e["hidden"], 2)):
            continue
        upcoming = [k for d in ("short", "unallocated", "hidden")
                    for k in e["months"][d] if k >= as_of]
        projects.append({
            "project_id": pid, "project_name": pname(pid),
            **{k: (P.projects.get(pid) or {}).get(k)
               for k in ("project_type", "clinical_phase", "status")},
            "short_fte_months": fc(e["short"]),
            "unallocated_fte_months": fc(e["unallocated"]),
            "over_fte_months": fc(e["over"]),
            "hidden_shortfall_fte_months": round(e["hidden"], 2),
            "missing_share_of_demand": round(
                (e["short"] + e["unallocated"] + e["hidden"] * prap_io.CENTS) / e["demand"], 3)
            if e["demand"] else None,
            "runs": {d: [span(r) for r in runs_of(ks)] for d, ks in e["months"].items()},
            "next_short_month": k_to_iso(min(upcoming)) if upcoming else None,
        })
    projects.sort(key=lambda x: -(x["short_fte_months"] + x["unallocated_fte_months"]
                                  + x["hidden_shortfall_fte_months"]
                                  + x["over_fte_months"] * 0.5))

    # ---- people, ranked by how far and how long over ------------------------------
    over_by = defaultdict(list)
    for r in over_pm:
        over_by[r["sid"]].append(r)
    people_over = []
    for sid, rs in over_by.items():
        peak = max(rs, key=lambda r: r["fte"])
        mix = sorted(contrib[(sid, peak["k"])].items(), key=lambda kv: -kv[1])
        people_over.append({
            "person_id": sid, "person_name": sname(sid),
            "department": (P.people.get(sid) or {}).get("department"),
            "primary_role": (P.people.get(sid) or {}).get("primary_role"),
            "over_months": len(rs),
            "excess_fte_months": fc(sum(cents(r["fte"]) - cents(P.OVER) for r in rs)),
            "peak_fte": peak["fte"], "peak_month": k_to_iso(peak["k"]),
            "runs": [span(r) for r in runs_of([r["k"] for r in rs])],
            "projects_at_peak": [{"project_id": p, "fte": round(v, 2)} for p, v in mix],
            "upcoming": any(r["k"] >= as_of for r in rs),
        })
    people_over.sort(key=lambda x: -x["excess_fte_months"])

    kp = defaultdict(list)
    for pid, k, sid, share in key_person:
        kp[(pid, sid)].append((k, share))
    key_people = sorted(
        ({"project_id": pid, "project_name": pname(pid), "person_id": sid,
          "person_name": sname(sid), "months": len(v),
          "runs": [span(r) for r in runs_of([k for k, _ in v])],
          "max_share": round(max(s for _, s in v), 2)}
         for (pid, sid), v in kp.items()),
        key=lambda x: -x["months"])

    under = [{"person_id": sid, "person_name": sname(sid),
              "department": (P.people.get(sid) or {}).get("department"),
              "months": run[2], "span": span(run)} for sid, run in under_runs]
    under.sort(key=lambda x: -x["months"])

    # ---- forward look --------------------------------------------------------------
    ahead = [k for k in P.months if as_of <= k < as_of + LOOKAHEAD]
    monthly = []
    for k in P.months:
        rows = [r for r in P.pm if r["k"] == k]
        monthly.append({
            "month": k_to_iso(k),
            "demand_fte": fc(sum(cents(r["demand"]) for r in rows)),
            "staffed_fte": fc(sum(cents(r["staffed"]) for r in rows)),
            "short_fte": fc(sum(-cents(r["gap"]) for r in rows if r["dir"] == "short")),
            "unallocated_fte": fc(sum(cents(r["demand"]) for r in rows
                                      if r["dir"] == "unallocated")),
            "hidden_shortfall_fte": round(sum(v for (_p, kk), v in hidden.items()
                                              if kk == k), 2),
            "people_over": sum(1 for r in over_pm if r["k"] == k),
            "people_staffed": sum(1 for r in staffed_pm if r["k"] == k),
            "projects_running": len({r["pid"] for r in rows}),
        })
    mon = {m["month"]: m for m in monthly}
    fwd = [mon[k_to_iso(k)] for k in ahead]
    before = [mon[k_to_iso(k)] for k in P.months if as_of - LOOKAHEAD <= k < as_of]
    avg = lambda xs, f: round(sum(x[f] for x in xs) / len(xs), 2) if xs else None
    forward = {
        "as_of": k_to_iso(as_of), "months": len(fwd),
        "peak_demand_month": max(fwd, key=lambda m: m["demand_fte"])["month"] if fwd else None,
        "peak_people_over_month": max(fwd, key=lambda m: m["people_over"])["month"]
        if fwd and any(m["people_over"] for m in fwd) else None,
        "avg_demand_next": avg(fwd, "demand_fte"),
        "avg_demand_previous": avg(before, "demand_fte"),
        "missing_fte_next": round(sum(m["short_fte"] + m["unallocated_fte"] for m in fwd), 2),
        "hidden_shortfall_fte_next": round(sum(m["hidden_shortfall_fte"] for m in fwd), 2),
        "over_person_months_next": sum(m["people_over"] for m in fwd),
        "starting_soon": sorted({p["project_id"] for p in projects
                                 if p["next_short_month"]
                                 and iso_to_k(p["next_short_month"]) < as_of + LOOKAHEAD}),
    }

    # ---- candidate moves -------------------------------------------------------------
    did = defaultdict(set)                      # sid -> pids worked on
    for L in P.lines:
        if L["fte"] > 0:
            did[L["sid"]].add(L["pid"])
    def kind(pid):
        p = P.projects.get(pid) or {}
        return (p.get("project_type"), p.get("clinical_phase"))
    roles_on = defaultdict(set)
    for L in P.lines:
        roles_on[L["pid"]].add(L["role"])
    everyone = sorted({r["sid"] for r in P.psm})

    def headroom(sid, ks):
        cap = (P.people.get(sid) or {}).get("capacity_fte")
        try:
            cap = float(cap)
        except (TypeError, ValueError):
            cap = 1.0
        limit = min(cap, P.OVER)
        # No PersonMonth row means not employed that month - not free.
        return min((limit - fte_of[(sid, k)]) if (sid, k) in fte_of else -1.0
                   for k in ks)

    def candidates(pid, ks, exclude=()):
        out = []
        for sid in everyone:
            if sid in exclude:
                continue
            room = headroom(sid, ks)
            if room < MIN_HEADROOM - 1e-9:
                continue
            fit = (3 if pid in did[sid] else
                   2 if any(kind(q) == kind(pid) for q in did[sid]) else
                   1 if any(kind(q)[0] == kind(pid)[0] for q in did[sid]) else 0)
            out.append({"person_id": sid, "person_name": sname(sid),
                        "primary_role": (P.people.get(sid) or {}).get("primary_role"),
                        "department": (P.people.get(sid) or {}).get("department"),
                        "min_headroom_fte": round(room, 2), "fit": fit})
        out.sort(key=lambda c: (-c["fit"], -c["min_headroom_fte"]))
        return out[:5]

    # The NEXT stretch of each problem, not all of it: a shortfall that runs for three
    # years has nobody free for all three, and "nobody" is not an answer anyone can use.
    # The first LOOKAHEAD months from where it next occurs is what can be acted on now.
    near = lambda ks: [k for k in ks if k < min(ks) + LOOKAHEAD] if ks else ks
    moves = []
    for p in projects:
        pid = p["project_id"]
        ks = near(sorted({r["k"] for r in P.pm if r["pid"] == pid
                          and r["dir"] in ("short", "unallocated") and r["k"] >= as_of}
                         | {k for (q, k), v in hidden.items()
                            if q == pid and k >= as_of and round(v, 2) > 0}))
        if not ks:
            continue
        kset = set(ks)
        moves.append({"for": "project", "project_id": pid,
                      "months": [span(r) for r in runs_of(ks)],
                      "missing_fte_months": round(fc(sum(
                          -cents(r["gap"]) for r in P.pm if r["pid"] == pid
                          and r["k"] in kset and r["dir"] in ("short", "unallocated")))
                          + sum(v for (q, k), v in hidden.items() if q == pid and k in kset), 2),
                      "roles_on_project": sorted(x for x in roles_on[p["project_id"]] if x),
                      "candidates": candidates(p["project_id"], ks)})
    for o in people_over:
        ks = near([r["k"] for r in over_by[o["person_id"]] if r["k"] >= as_of])
        if not ks:
            continue
        biggest = o["projects_at_peak"][0]["project_id"] if o["projects_at_peak"] else None
        moves.append({"for": "person", "person_id": o["person_id"],
                      "months": [span(r) for r in runs_of(ks)],
                      "offload_from_project": biggest,
                      "candidates": candidates(biggest, ks, exclude={o["person_id"]})
                      if biggest else []})

    return {
        "scorecard_version": SCORING["version"],
        "source": {"file": P.source, "kind": P.kind, "filters": P.filters,
                   "horizon": f"{k_to_iso(P.months[0])} to {k_to_iso(P.months[-1])}"
                   if P.months else None, "months": len(P.months),
                   "projects": len({r["pid"] for r in P.pm}),
                   "people": len({r["sid"] for r in P.psm}),
                   "thresholds": {"over_fte": P.OVER, "under_fte": P.UNDER,
                                  "under_min_months": P.MINM}},
        "score": score, "band": band(score), "components": comps,
        "projects_off_standard": projects[:top], "projects_off_standard_total": len(projects),
        "people_over_allocated": people_over[:top], "people_over_total": len(people_over),
        "key_person_dependencies": key_people[:top],
        "key_person_total": len(key_people),
        "under_allocation_runs": under[:top], "under_runs_total": len(under),
        "forward": forward, "monthly": monthly,
        "candidate_moves": moves[:top], "candidate_moves_total": len(moves),
        "notes": [
            "Every figure here is computed by tools/prap_analyze.py from the file named "
            "in source; nothing is estimated. Interpret it, do not recompute it.",
            "Candidate moves are indicative: the engine divides a project's demand among "
            "whoever is on it, so adding a person changes every share. Confirm a move by "
            "editing the source plan and running `prap_analyze.py compare`.",
        ] + ([f"This file was exported with a filter on ({P.filters}). The scorecard "
              "covers only what was in view; staffed and gap figures are still whole "
              "projects."] if P.filters and not P.filters.lower().startswith("none")
             else []),
    }


# ------------------------------------------------------------------------- privacy
def pseudonymise(P, key_path):
    depts = {}
    rows = [("person_id", "person_name", "department", "department_code")]
    for sid, p in sorted(P.people.items()):
        d = p.get("department")
        code = depts.setdefault(d, f"D{len(depts) + 1:02d}") if d else None
        rows.append((sid, p.get("person_name"), d, code))
        p["person_name"] = None
        p["department"] = code
    Path(key_path).parent.mkdir(parents=True, exist_ok=True)
    Path(key_path).write_text("\n".join(",".join("" if v is None else str(v) for v in r)
                                        for r in rows) + "\n", encoding="utf-8")


# ------------------------------------------------------------------------ rendering
def md_score(A):
    s = A["source"]
    out = [f"# Resource plan scorecard — {A['score']} / 100 ({A['band']})", "",
           f"Source: `{s['file']}` ({s['kind']}), {s['horizon']}, {s['projects']} projects, "
           f"{s['people']} people. Filters: {s['filters']}. "
           f"Thresholds: over {s['thresholds']['over_fte']}, under "
           f"{s['thresholds']['under_fte']} for {s['thresholds']['under_min_months']}+ months. "
           f"Scorecard v{A['scorecard_version']}.", "",
           "Each component scores weight × 0.5^(ratio ÷ half at): full marks at 0, half "
           "at the stated ratio.", "",
           "| Component | Points | Ratio | Half at | Inputs |", "|---|---|---|---|---|"]
    for name, c in A["components"].items():
        ins = ", ".join(f"{k} {v}" for k, v in c["inputs"].items())
        out.append(f"| {name} | {c['points']} / {c['of']} | {c['ratio']} | "
                   f"{c['half_at']} | {ins} |")
    out += ["", f"## Projects short of resource or off their standard "
            f"({A['projects_off_standard_total']})", ""]
    if A["projects_off_standard"]:
        out += ["FTE-months. *Hidden* is the overload of this project's people above the "
                "ceiling, attributed back to it: the engine shares a project's demand among "
                "whoever is on it, so a project with too few people shows up as its people "
                "going over, never as a gap.", "",
                "| Project | Short | Unallocated | Hidden | Over standard | Missing share "
                "| Next short month |", "|---|---|---|---|---|---|---|"]
        for p in A["projects_off_standard"]:
            out.append(f"| {p['project_id']} {p['project_name'] or ''} | "
                       f"{p['short_fte_months']} | {p['unallocated_fte_months']} | "
                       f"{p['hidden_shortfall_fte_months']} | "
                       f"{p['over_fte_months']} | {p['missing_share_of_demand']} | "
                       f"{p['next_short_month'] or '-'} |")
    else:
        out.append("None — every project-month is on its standard.")
    out += ["", f"## People over the ceiling ({A['people_over_total']})", ""]
    if A["people_over_allocated"]:
        out += ["| Person | Months over | Excess FTE-months | Peak | Projects at peak |",
                "|---|---|---|---|---|"]
        for o in A["people_over_allocated"]:
            mix = ", ".join(f"{x['project_id']} {x['fte']}" for x in o["projects_at_peak"][:4])
            out.append(f"| {o['person_id']} {o['person_name'] or ''} | {o['over_months']} | "
                       f"{o['excess_fte_months']} | {o['peak_fte']} in {o['peak_month']} | "
                       f"{mix} |")
    else:
        out.append("None.")
    f = A["forward"]
    out += ["", f"## The next {f['months']} months from {f['as_of']}", "",
            f"- Average demand {f['avg_demand_next']} FTE a month "
            f"(previous {LOOKAHEAD} months: "
            f"{'not in the file' if f['avg_demand_previous'] is None else f['avg_demand_previous']}"
            f"); peak in "
            f"{f['peak_demand_month']}.",
            f"- Missing (short + unallocated): {f['missing_fte_next']} FTE-months; hidden "
            f"in overload: {f['hidden_shortfall_fte_next']} FTE-months; "
            f"person-months over the ceiling: {f['over_person_months_next']}"
            + (f", most in {f['peak_people_over_month']}" if f['peak_people_over_month'] else "")
            + ".",
            f"- Projects whose shortfall starts within the window: "
            f"{', '.join(f['starting_soon']) or 'none'}."]
    out += ["", f"## Candidate moves ({A['candidate_moves_total']}) — the next {LOOKAHEAD} "
            "months of each problem; indicative, confirm with `compare`", ""]
    for m in A["candidate_moves"]:
        who = ", ".join(f"{c['person_id']} (room {c['min_headroom_fte']}, fit {c['fit']})"
                        for c in m["candidates"]) or "nobody has room in those months"
        if m["for"] == "project":
            out.append(f"- **{m['project_id']}** missing {m['missing_fte_months']} FTE-months "
                       f"in {', '.join(m['months'])}: {who}")
        else:
            out.append(f"- **{m['person_id']}** over in {', '.join(m['months'])}; offload "
                       f"from {m['offload_from_project']} to: {who}")
    out += ["", "Fit: 3 already on the project, 2 has worked on the same type and phase, "
            "1 same type, 0 none.", ""] + [f"> {n}" for n in A["notes"]]
    return "\n".join(out) + "\n"


def compare(A, B):
    d = {"score": {"before": A["score"], "after": B["score"],
                   "change": round(B["score"] - A["score"], 1),
                   "band": [A["band"], B["band"]]},
         "components": {k: {"before": A["components"][k]["points"],
                            "after": B["components"][k]["points"],
                            "change": round(B["components"][k]["points"]
                                            - A["components"][k]["points"], 1)}
                        for k in A["components"]}}
    pa = {p["project_id"]: p for p in A["_all_projects"]}
    pb = {p["project_id"]: p for p in B["_all_projects"]}
    rows = []
    for pid in sorted(set(pa) | set(pb)):
        a, b = pa.get(pid, {}), pb.get(pid, {})
        miss = lambda x: round(x.get("short_fte_months", 0) + x.get("unallocated_fte_months", 0)
                               + x.get("hidden_shortfall_fte_months", 0), 2)
        if miss(a) != miss(b) or a.get("over_fte_months", 0) != b.get("over_fte_months", 0):
            row = {"project_id": pid, "missing_before": miss(a), "missing_after": miss(b),
                   "over_before": a.get("over_fte_months", 0),
                   "over_after": b.get("over_fte_months", 0)}
            for part in ("short", "unallocated", "hidden_shortfall"):
                row[f"{part}_before"] = a.get(f"{part}_fte_months", 0)
                row[f"{part}_after"] = b.get(f"{part}_fte_months", 0)
            rows.append(row)
    d["projects_changed"] = rows
    oa = {o["person_id"]: o for o in A["_all_people_over"]}
    ob = {o["person_id"]: o for o in B["_all_people_over"]}
    d["people_changed"] = [
        {"person_id": sid, "over_months_before": oa.get(sid, {}).get("over_months", 0),
         "over_months_after": ob.get(sid, {}).get("over_months", 0)}
        for sid in sorted(set(oa) | set(ob))
        if oa.get(sid, {}).get("over_months", 0) != ob.get(sid, {}).get("over_months", 0)]
    d["forward"] = {"missing_fte_next": [A["forward"]["missing_fte_next"],
                                         B["forward"]["missing_fte_next"]],
                    "hidden_shortfall_fte_next": [A["forward"]["hidden_shortfall_fte_next"],
                                                  B["forward"]["hidden_shortfall_fte_next"]],
                    "over_person_months_next": [A["forward"]["over_person_months_next"],
                                                B["forward"]["over_person_months_next"]]}
    return d


def md_compare(D, a_src, b_src, horizon_):
    s = D["score"]
    out = [f"# Scenario comparison — {s['before']} → {s['after']} ({s['change']:+})", "",
           f"Before: `{a_src}`. After: `{b_src}`. Compared over {horizon_}, the months "
           "both cover.", "", "| Component | Before | After | Change |", "|---|---|---|---|"]
    for k, c in D["components"].items():
        out.append(f"| {k} | {c['before']} | {c['after']} | {c['change']:+} |")
    out += ["", "## Projects whose shortfall or excess changed", "",
            "Missing = short + unallocated + hidden in overload, in FTE-months.", ""]
    out += ([f"- {r['project_id']}: missing {r['missing_before']} → {r['missing_after']} "
             f"FTE-months (short {r['short_before']} → {r['short_after']}, unallocated "
             f"{r['unallocated_before']} → {r['unallocated_after']}, hidden "
             f"{r['hidden_shortfall_before']} → {r['hidden_shortfall_after']}), over standard "
             f"{r['over_before']} → {r['over_after']}"
             for r in D["projects_changed"]] or ["None."])
    out += ["", "## People whose months over the ceiling changed", ""]
    out += ([f"- {r['person_id']}: {r['over_months_before']} → {r['over_months_after']} months"
             for r in D["people_changed"]] or ["None."])
    f = D["forward"]
    out += ["", f"Next {LOOKAHEAD} months: missing {f['missing_fte_next'][0]} → "
            f"{f['missing_fte_next'][1]} FTE-months; hidden in overload "
            f"{f['hidden_shortfall_fte_next'][0]} → {f['hidden_shortfall_fte_next'][1]}; "
            f"person-months over "
            f"{f['over_person_months_next'][0]} → {f['over_person_months_next'][1]}.", ""]
    return "\n".join(out)


# ------------------------------------------------------------------------------- cli
def full(P, as_of):
    """analyse(), keeping the unabridged project and people lists for compare()."""
    A = analyse(P, as_of)
    every = analyse(P, as_of, top=10 ** 9)
    A["_all_projects"] = every["projects_off_standard"]
    A["_all_people_over"] = every["people_over_allocated"]
    return A


def write(out_dir, name, data, md):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{name}.json").write_text(json.dumps(data, indent=1, ensure_ascii=False,
                                                 default=str), encoding="utf-8")
    (out / f"{name}.md").write_text(md, encoding="utf-8")
    return out / f"{name}.json", out / f"{name}.md"


def as_of_key(s):
    if s:
        return iso_to_k(s)
    t = date.today()
    return t.year * 12 + t.month - 1


def main(argv=None):
    ap = argparse.ArgumentParser(prog="prap_analyze", description=__doc__.split("\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("score", "compare"):
        p = sub.add_parser(name)
        p.add_argument("files", nargs=1 if name == "score" else 2)
        p.add_argument("--as-of", metavar="YYYY-MM",
                       help="the month the forward look starts from (default: this month)")
        p.add_argument("--from", dest="frm", metavar="YYYY-MM")
        p.add_argument("--to", metavar="YYYY-MM")
        p.add_argument("--out", default=".", help="folder for the .json and .md")
        p.add_argument("--pseudonymise", action="store_true",
                       help="drop person names and code departments; the key goes to a "
                            "separate local file that is not for sharing")
        p.add_argument("--quiet", action="store_true")
    a = ap.parse_args(argv)
    as_of = as_of_key(a.as_of)
    plans = [load(f) for f in a.files]
    for P in plans:
        horizon(P, a.frm, a.to)
    if len(plans) == 2:                         # compare over the months both cover
        both = sorted(set(plans[0].months) & set(plans[1].months))
        if not both:
            raise SystemExit("error: the two plans share no month")
        for P in plans:
            P.restrict(both)
    if a.pseudonymise:
        for i, P in enumerate(plans):
            pseudonymise(P, Path(a.out) / f"pseudonym_key{'' if i == 0 else '_after'}.csv")

    if a.cmd == "score":
        A = analyse(plans[0], as_of)
        md = md_score(A)
        paths = write(a.out, "scorecard", A, md)
    else:
        A, B = full(plans[0], as_of), full(plans[1], as_of)
        D = compare(A, B)
        for X in (A, B):
            X.pop("_all_projects"), X.pop("_all_people_over")
        D["before"], D["after"] = A, B
        hz = A["source"]["horizon"]
        md = md_compare(D, A["source"]["file"], B["source"]["file"], hz)
        paths = write(a.out, "comparison", D, md)
    if not a.quiet:
        print(md)
    print(f"written: {paths[0]}  {paths[1]}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
