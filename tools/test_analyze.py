"""The scorecard an AI is handed must be the same whichever file it was made from.

tools/prap_analyze.py scores a plan from either the calculated export or the source
plan itself. Those are two routes to one set of figures - the browser engine through
the export, the Python reference through the source - and a scorecard that differed
between them would make a what-if comparison measure the route instead of the change.

What is checked:

  * THE TWO ROUTES AGREE: the export the application writes and the source plan it was
    written from give the same scorecard, component for component and row for row.
  * The things the export now carries reach the scorecard: a project nobody is on is
    listed as unallocated with exactly its demand, and the free months of a person on
    nothing are offered as room.
  * It is reproducible - the same file gives the same result twice.
  * `compare` measures a change: putting the missing staff back raises demand coverage
    and clears the project.
  * --pseudonymise leaves no person name in what would be shared, and keeps the key apart.
  * An export from before R-55 is refused with a reason, not scored as if nothing were
    short.

    python tools/test_analyze.py
"""

import calendar
import json
import pathlib
import subprocess
import sys
import tempfile
from datetime import date

from openpyxl import Workbook
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parents[1]
APP = (ROOT / "app" / "PRAP.html").as_uri()
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
FIX = ROOT / "templates" / "PRAP_SourceData_Dummy_10x10_v1.13.xlsx"
TMP = pathlib.Path(tempfile.mkdtemp(prefix="prap_analyze_"))

sys.path.insert(0, str(ROOT / "tools"))
import prap_analyze as PA                                              # noqa: E402
import prap_io                                                         # noqa: E402

fails = []


def check(ok, label, detail=""):
    print(f"  {'ok  ' if ok else 'FAIL'} {label}{'   ' + detail if detail else ''}")
    if not ok:
        fails.append(label)


def strip(A):
    """Everything but the provenance, which names the file and so must differ."""
    return {k: v for k, v in A.items() if k != "source"}


# ---- a plan with something wrong in every direction ----------------------------------
src = prap_io.read_xlsx(FIX)
gapped = {pid for (pid, _k) in prap_io.calculate(prap_io.Model(prap_io.read_xlsx(FIX)))["proj_gap"]}
gone = next(p["project_id"] for p in src["Project"]
            if str(p.get("status") or "") != "Completed" and p["project_id"] not in gapped
            and any(a["project_id"] == p["project_id"] for a in src["Assignment"]))
src["Assignment"] = [a for a in src["Assignment"] if a["project_id"] != gone]
keep = {a["assignment_id"] for a in src["Assignment"]}
for sh in ("PersonPeriodWeight", "MonthlyEstimate"):
    src[sh] = [r for r in src.get(sh, []) if not r.get("assignment_id") or r["assignment_id"] in keep]
t = date.today()
y, m = divmod(t.month - 1 + 3, 12)
src["Person"][-1]["employment_end"] = date(t.year + y, m + 1,
                                           calendar.monthrange(t.year + y, m + 1)[1])
broken = TMP / "broken.xlsx"
prap_io.write_xlsx(src, broken)
AS_OF = f"{t.year}-{t.month:02d}"

print("tools/prap_analyze.py — the scorecard, by both routes")
with sync_playwright() as pw:
    browser = pw.chromium.launch(executable_path=CHROME, downloads_path=str(TMP))
    pg = browser.new_context(accept_downloads=True).new_page()
    pg.set_viewport_size({"width": 1500, "height": 950})
    pg.set_default_timeout(25000)
    errors = []
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.on("dialog", lambda d: d.accept())
    pg.goto(APP)
    pg.wait_for_timeout(400)
    pg.set_input_files("#picker", str(broken))
    pg.wait_for_timeout(3500)
    pg.click("#exportBtn")
    pg.wait_for_timeout(300)
    with pg.expect_download() as dl:
        pg.click("#exportCalcBtn")
    exported = TMP / "calc.xlsx"
    dl.value.save_as(exported)
    check(not errors, "the application exported without an error", "; ".join(errors[:2]))
    browser.close()

as_of = PA.iso_to_k(AS_OF)
Pe = PA.load(exported)
Ps = PA.load(broken)
Ps.restrict(Pe.months)
Ae, As = PA.analyse(Pe, as_of), PA.analyse(Ps, as_of)
check(Pe.kind == "calculated export" and Ps.kind == "source plan",
      "one file is read as the export and the other as the source plan")
diff = [k for k in strip(Ae) if json.dumps(Ae[k], default=str) != json.dumps(As[k], default=str)]
check(not diff,
      "THE TWO ROUTES AGREE — the application's export and the source plan give the same "
      "scorecard, component for component and row for row",
      f"score {Ae['score']} / {As['score']}" + (f"; differ in {', '.join(diff)}" if diff else ""))
check(PA.analyse(PA.load(exported), as_of) == Ae, "and the same file scores the same twice")

row = next((p for p in Ae["projects_off_standard"] if p["project_id"] == gone), None)
from openpyxl import load_workbook                                       # noqa: E402
pm = PA._sheet(load_workbook(exported, read_only=True), "ProjectMonth")
want = round(sum(r["demand_fte"] for r in pm if r["project_id"] == gone), 2)
check(row and row["unallocated_fte_months"] == want and want > 0,
      "A PROJECT NOBODY IS ON IS LISTED AS UNALLOCATED, with exactly its demand",
      f"{gone}: {row and row['unallocated_fte_months']} of {want} FTE-months")
idle = {r["sid"] for r in Pe.psm if r["fte"] == 0}
offered = {c["person_id"] for mv in Ae["candidate_moves"] for c in mv["candidates"]}
check(idle and offered,
      "the months somebody is on nothing reach the scorecard as room to offer",
      f"{len(idle)} person(s) with a free month; {len(offered)} offered as candidates")
leaver = src["Person"][-1]["person_id"]
late = [k for k in Pe.months if k > as_of + 3]
check(late and not any(r["sid"] == leaver and r["k"] in late for r in Pe.psm if r["fte"] == 0),
      "and somebody who has left is not")

every = PA.analyse(Pe, as_of, top=10 ** 9)
hid = sum(p["hidden_shortfall_fte_months"] for p in every["projects_off_standard"])
exc = sum(o["excess_fte_months"] for o in every["people_over_allocated"])
check(exc > 0 and abs(hid - exc) < 0.01 * max(1, len(every["projects_off_standard"])),
      "THE SHORTFALL HIDDEN IN OVERLOAD IS ACCOUNTED FOR — what the projects are given as "
      "hidden adds up to what their people carry above the ceiling, no more and no less",
      f"{hid:.2f} attributed to projects, {exc:.2f} excess on people")

comps = Ae["components"]
check(abs(sum(c["points"] for c in comps.values()) - Ae["score"]) < 0.05
      and sum(c["of"] for c in comps.values()) == 100
      and all(0 <= c["points"] <= c["of"] for c in comps.values()),
      "the score is the sum of its components, which are weighted to 100",
      " · ".join(f"{k} {c['points']}/{c['of']}" for k, c in comps.items()))

# ---- compare: put the staff back, and the scorecard says so ------------------------
out = TMP / "cmp"
r = subprocess.run([sys.executable, "-W", "error", str(ROOT / "tools" / "prap_analyze.py"),
                    "compare", str(broken), str(FIX), "--as-of", AS_OF, "--out", str(out),
                    "--quiet"], capture_output=True, text=True)
D = json.loads((out / "comparison.json").read_text()) if r.returncode == 0 else {}
ch = next((x for x in D.get("projects_changed", []) if x["project_id"] == gone), None)
check(r.returncode == 0 and D["components"]["demand_coverage"]["change"] > 0
      and ch and ch["unallocated_before"] > 0 and ch["unallocated_after"] == 0
      and ch["missing_after"] < ch["missing_before"],
      "COMPARE MEASURES A CHANGE — putting the staff back raises demand coverage and "
      "clears the project's unallocated months",
      (r.stderr.strip().splitlines() or [""])[-1] if r.returncode else
      f"coverage {D['components']['demand_coverage']['change']:+}; {gone} unallocated "
      f"{ch and ch['unallocated_before']} -> {ch and ch['unallocated_after']}, missing in all "
      f"{ch and ch['missing_before']} -> {ch and ch['missing_after']}")
check((out / "comparison.md").exists(), "and writes the comparison for a person as well")

# A plan already far over the ceiling must still register relief. A score that runs in
# a straight line to a floor goes flat past it, and then taking a person from 21 months
# over to 8 scores the same as doing nothing - which is how this was found.
base = prap_io.read_xlsx(FIX)
over_pid = "PRJ-003"
spare = min(base["Person"], key=lambda p: float(p.get("capacity_fte") or 1))["person_id"]
role = next(a["role_name"] for a in base["Assignment"] if a["project_id"] == over_pid)
base["Assignment"].append({"assignment_id": "ASG-900", "person_id": spare,
                           "project_id": over_pid, "role_name": role,
                           "assign_start_date": date(t.year, t.month, 1),
                           "assign_end_date": None, "person_weight": 0.3,
                           "estimation_type": "automatic"})
relieved = TMP / "relieved.xlsx"
prap_io.write_xlsx(base, relieved)
Pa, Pb = PA.load(FIX), PA.load(relieved)
both = sorted(set(Pa.months) & set(Pb.months))
Pa.restrict(both), Pb.restrict(both)
oa = PA.analyse(Pa, as_of)["components"]["over_allocation"]
ob = PA.analyse(Pb, as_of)["components"]["over_allocation"]
check(oa["ratio"] > 2 * oa["half_at"] and ob["points"] > oa["points"],
      "A PLAN FAR OVER THE CEILING STILL REGISTERS RELIEF — the score has no floor to go "
      "flat against",
      f"ratio {oa['ratio']} (half at {oa['half_at']}); points {oa['points']} -> {ob['points']}")

# ---- privacy ------------------------------------------------------------------------
out = TMP / "anon"
r = subprocess.run([sys.executable, "-W", "error", str(ROOT / "tools" / "prap_analyze.py"),
                    "score", str(exported), "--as-of", AS_OF, "--out", str(out),
                    "--pseudonymise", "--quiet"], capture_output=True, text=True)
names = [p["person_name"] for p in src["Person"] if p.get("person_name")]
shared = (out / "scorecard.json").read_text() + (out / "scorecard.md").read_text() \
    if r.returncode == 0 else ""
key = (out / "pseudonym_key.csv").read_text() if (out / "pseudonym_key.csv").exists() else ""
check(r.returncode == 0 and shared and not any(n in shared for n in names)
      and all(n in key for n in names),
      "--PSEUDONYMISE LEAVES NO PERSON NAME IN WHAT WOULD BE SHARED, and keeps the key apart",
      f"{len(names)} names checked" if r.returncode == 0 else r.stderr[-200:])
depts = {p.get("department") for p in src["Person"] if p.get("department")}
check(shared and depts and not any(f'"{d}"' in shared for d in depts),
      "and departments are coded", f"{len(depts)} department(s)")

# ---- an export from before R-55 -----------------------------------------------------
old = TMP / "old.xlsx"
wb = Workbook()
wb.active.title = "00_ReadMe"
ws = wb.create_sheet("ProjectMonth")
ws.append(["month", "month_iso", "project_id", "fte"])
ws.append(["Jan 2026", "2026-01", "PRJ-001", 1.0])
wb.save(old)
r = subprocess.run([sys.executable, str(ROOT / "tools" / "prap_analyze.py"), "score",
                    str(old), "--out", str(TMP / "old")], capture_output=True, text=True)
check(r.returncode != 0 and "R-55" in r.stderr,
      "AN EXPORT FROM BEFORE R-55 IS REFUSED WITH A REASON — not scored as if nothing "
      "were short", r.stderr.strip()[:110])

print(f"\nFAILURES: {'none' if not fails else len(fails)}")
for f in fails:
    print(f"  FAILED  {f}")
sys.exit(1 if fails else 0)
