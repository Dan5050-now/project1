#!/usr/bin/env python3
"""Shared openpyxl styling for the TEA controlled workbooks.

Extracted from the workbooks already in the repository so every document keeps
one visual identity. The original build scripts lived in a scratchpad and did
not survive the session that wrote them; this module and the build scripts that
import it are committed so the documents stay reproducible.

Layout note that is easy to get wrong: Excel auto-fits row height for wrapped
text only in UNMERGED cells. Any table that merges cells must set its row
heights explicitly or the text is clipped on open, and no check in this
repository would catch it. Use write_rows(), which does it.
"""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Palette
INK = "16201C"          # body text
ACCENT = "1D4A3A"       # headings, table headers
MUTE = "6B7770"         # subtitles, de-emphasised text
PAPER = "F6F7F5"        # zebra banding
WASH = "E3EDE7"         # section wash
INPUT_FILL = "FFF9D6"   # reviewer-editable cells
AMBER_LT = "FBF0DC"     # caution band

# Severity / classification colours
CRITICAL = "8C2F39"
MAJOR = "B06A1F"
MINOR = "4A6B8A"
INFO = "6B7770"

FACE = "Arial"

# The three portability classes, used across TEA-BND-001 and the AI guides.
CLASS_COLOUR = {
    "C": ACCENT,        # COMMON — built here, do not edit in the company
    "A": MAJOR,         # APPLY  — the company's Claude Code produces it
    "K": MINOR,         # KNOW   — context the company's Claude Code is given
}

# Excel auto-fits wrapped text only in unmerged cells; merged rows need a height.
CHARS_PER_WIDTH_UNIT = 0.90   # pessimistic: word wrap breaks lines early
POINTS_PER_LINE = 12.5


def sheet_setup(ws, widths, gridlines=False):
    """widths: dict of column letter -> width, or a list in column order."""
    if isinstance(widths, (list, tuple)):
        widths = {get_column_letter(i): w for i, w in enumerate(widths, 1)}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.sheet_view.showGridLines = gridlines
    return len(widths)


def title_style(ws, text, ncols, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, text)
    c.font = Font(name=FACE, size=13, bold=True, color=ACCENT)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        c = ws.cell(2, 1, subtitle)
        c.font = Font(name=FACE, size=9, color=MUTE)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 30


def header_row(ws, labels, row, spans=None):
    spans = spans or [1] * len(labels)
    col = 1
    for label, span in zip(labels, spans):
        c = ws.cell(row, col, label)
        c.font = Font(name=FACE, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        for k in range(1, span):
            ws.cell(row, col + k).fill = PatternFill("solid", fgColor=ACCENT)
        if span > 1:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=row, end_column=col + span - 1)
        col += span
    ws.row_dimensions[row].height = 30
    return row + 1


def section_label(ws, text, row, note=None):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FACE, size=11, bold=True, color=ACCENT)
    row += 1
    if note:
        c = ws.cell(row, 1, note)
        c.font = Font(name=FACE, size=9, italic=True, color=MUTE)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        row += 1
    return row


def set_row_height(ws, row, values, spans):
    """Compute a height for a row whose cells are merged. No-op when unmerged."""
    if all(s == 1 for s in spans):
        return
    widths = [ws.column_dimensions[get_column_letter(c)].width or 8.43
              for c in range(1, sum(spans) + 1)]
    lines, col = 1, 0
    for val, span in zip(values, spans):
        usable = sum(widths[col:col + span]) * CHARS_PER_WIDTH_UNIT
        lines = max(lines, -(-len(str(val or "")) // max(1, int(usable))))
        col += span
    # One spare line: over-sizing a row costs nothing, clipping loses content.
    ws.row_dimensions[row].height = min(409, (lines + 1) * POINTS_PER_LINE + 4)


def write_rows(ws, rows, start, bold_cols=(), colour_cols=None, zebra=True,
               spans=None):
    """Write data rows.

    bold_cols    0-based logical column indices to embolden.
    colour_cols  0-based index -> {cell value: hex colour}; those cells are bold.
    spans        logical column -> width in physical columns, so a secondary
                 table can use its own layout on a sheet whose widths were set
                 for the main table.
    """
    colour_cols = colour_cols or {}
    spans = spans or [1] * max(len(r) for r in rows)
    r = start
    for n, row in enumerate(rows):
        band = zebra and n % 2 == 1
        col = 1
        for i, val in enumerate(row):
            c = ws.cell(r, col, val)
            colour = colour_cols[i].get(str(val), INK) if i in colour_cols else INK
            c.font = Font(name=FACE, size=9,
                          bold=(i in bold_cols or i in colour_cols), color=colour)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if spans[i] > 1:
                ws.merge_cells(start_row=r, start_column=col,
                               end_row=r, end_column=col + spans[i] - 1)
            if band:
                for k in range(spans[i]):
                    ws.cell(r, col + k).fill = PatternFill("solid", fgColor=PAPER)
            col += spans[i]
        set_row_height(ws, r, row, spans)
        r += 1
    return r


def note_block(ws, row, ncols, text, fill=AMBER_LT, bold=False):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, text)
    c.font = Font(name=FACE, size=9, bold=bold, color=INK)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    for i in range(1, ncols + 1):
        ws.cell(row, i).fill = PatternFill("solid", fgColor=fill)
    set_row_height(ws, row, [text], [ncols])
    return row + 1


def bullets(ws, row, ncols, items, indent=1):
    for item in items:
        ws.merge_cells(start_row=row, start_column=indent,
                       end_row=row, end_column=ncols)
        c = ws.cell(row, indent, f"•  {item}")
        c.font = Font(name=FACE, size=9, color=INK)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        set_row_height(ws, row, [item], [ncols - indent + 1])
        row += 1
    return row


def kv_rows(ws, row, pairs, key_width_cols=1, ncols=6):
    """Label/value rows, the shape the Cover sheets use."""
    for k, v in pairs:
        a = ws.cell(row, 1, k)
        a.font = Font(name=FACE, size=9, bold=True, color=ACCENT)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(row, 1 + key_width_cols, v)
        b.font = Font(name=FACE, size=9, color=INK)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        if ncols > 1 + key_width_cols:
            ws.merge_cells(start_row=row, start_column=1 + key_width_cols,
                           end_row=row, end_column=ncols)
        set_row_height(ws, row, [k, v], [key_width_cols, ncols - key_width_cols])
        row += 1
    return row


def find_row(ws, key_in_col_a):
    """Row whose column A holds this key, else the first free row.

    Keeps append-style edits idempotent across re-runs.
    """
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == key_in_col_a:
            return r
    return ws.max_row + 1
