"""Capture what V-36 looks like in the application, into output/.

Not a test - tools/test_unstaffed.py does that. This is so a reader can SEE the rule
without opening the application, and see the thing it is FOR: the same project, before
and after its first assignment.

The fixture is built here rather than committed, for the reason test_unstaffed.py gives:
the point is the DIFFERENCE between this and the delivered plan, and a stale copy would
stop being the same plan.

    python tools/shots_unstaffed.py
"""

import pathlib
import shutil
import tempfile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parents[1]
APP = (ROOT / "app" / "PRAP.html").as_uri()
DUMMY = ROOT / "templates" / "PRAP_SourceData_Dummy_10x10_v1.11.xlsx"
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)

NEW = "PRJ-099"


def make(dst, staffed=False):
    """The delivered plan plus one project cloned from PRJ-001 - same type, phase, scope
    and periods - with nobody on it unless asked for."""
    shutil.copy(DUMMY, dst)
    wb = load_workbook(dst)
    p, per, asg = wb["Project"], wb["ProjectPeriod"], wb["Assignment"]
    cols = [c.value for c in p[1]]
    row = [c.value for c in p[2]]
    row[cols.index("project_id")] = NEW
    row[cols.index("project_name")] = "ZZZ-900 Phase 1 (just created)"
    p.append(row)
    for r in per.iter_rows(min_row=2, values_only=True):
        if r[0] == "PRJ-001":
            nr = list(r)
            nr[0] = NEW
            per.append(nr)
    if staffed:
        ac = [c.value for c in asg[1]]
        a = {c: None for c in ac}
        # Blank dates on purpose: REQ-CAL-15 makes them the project's own, which is the
        # window V-36 predicted from.
        a.update({"assignment_id": "ASG-999", "person_id": "PSN-001",
                  "project_id": NEW, "role_name": "Lead data manager",
                  "person_weight": 1.0, "estimation_type": "automatic"})
        asg.append([a.get(c) for c in ac])
    wb.save(dst)


def load(pg, path):
    pg.goto(APP)
    pg.wait_for_timeout(250)
    pg.set_input_files("#picker", str(path))
    pg.wait_for_function("() => !document.getElementById('tabs').hidden", timeout=60000)
    pg.wait_for_timeout(1200)


def shot(pg, sel, name):
    pg.locator(sel).first.screenshot(path=str(OUT / name))
    print(f"  {name}")


tmp = pathlib.Path(tempfile.mkdtemp())
bare, staffed = tmp / "unstaffed.xlsx", tmp / "staffed.xlsx"
make(bare)
make(staffed, staffed=True)

with sync_playwright() as pw:
    browser = pw.chromium.launch(executable_path=CHROME)
    pg = browser.new_page(viewport={"width": 1600, "height": 1000},
                          device_scale_factor=2)

    print("output/ —")
    load(pg, bare)

    # 1. THE BANNER, which is the first thing a user sees on opening the plan.
    shot(pg, "#banner", "v36_1_banner.png")

    # 2. THE FULL REPORT, with the note in it. Scrolled to the V-36 row so the figure
    #    is legible rather than the dialog being shown whole and unreadable.
    pg.evaluate("() => { renderReport(S.model.findings); document.getElementById('report').showModal(); }")
    pg.wait_for_timeout(400)
    pg.evaluate("""() => {
      const tr = [...document.querySelectorAll('#report table tbody tr')]
        .find(r => r.innerText.includes('V-36'));
      if (tr) tr.scrollIntoView({block:'center'});
    }""")
    pg.wait_for_timeout(300)
    shot(pg, "#report", "v36_2_report.png")
    pg.keyboard.press("Escape")
    pg.wait_for_timeout(300)

    # 3 and 4. THE DEFECT ITSELF, and the same row once somebody is on it. The PANEL
    #    rather than the whole tab: a full-tab picture of this application is eleven
    #    thousand pixels tall and legible at no size anybody will view it at, which
    #    makes it decorative rather than evidence.
    def overall(pg, tag, colour):
        pg.click('nav button[data-tab="t-overall"]')
        pg.wait_for_timeout(900)
        marks = pg.evaluate(f"""() => {{
          const rows = [...document.querySelectorAll('#t-overall .grid-t tbody tr')];
          const hit = rows.find(r => r.innerText.includes('{NEW}'));
          if (hit) hit.style.outline = '3px solid {colour}';
          return [...document.querySelectorAll('#t-overall .panel')].map(p => {{
            const h = ((p.querySelector('h2') || {{}}).textContent || '').trim();
            const n = [...p.querySelectorAll('svg.chart [data-s]')].filter(e =>
              (e.getAttribute('data-s') || '').includes('{NEW}')
              || (e.getAttribute('data-s2') || '').includes('{NEW}')).length;
            return p.querySelector('svg.chart') ? [h.slice(0, 28), n] : null;
          }}).filter(Boolean);
        }}""")
        print(f"     marks on each chart ({tag}): "
              + ", ".join(f"{h} {n}" for h, n in marks))
        for starts, short in (("Resource by project", "rbp"), ("Project timeline", "tl")):
            h = pg.evaluate_handle(
                """(s) => [...document.querySelectorAll('#t-overall .panel')]
                     .find(e => ((e.querySelector('h2')||{}).textContent||'').startsWith(s))""",
                starts)
            h.as_element().screenshot(path=str(OUT / f"v36_{short}_{tag}.png"))
            print(f"  v36_{short}_{tag}.png")

    overall(pg, "before", "#c00")
    load(pg, staffed)
    overall(pg, "after", "#090")
    shot(pg, "#banner", "v36_5_banner_after.png")

    browser.close()

shutil.rmtree(tmp, ignore_errors=True)
