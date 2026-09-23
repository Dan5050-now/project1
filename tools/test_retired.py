"""MonthlyEstimate.edited_at is retired, and a workbook that still has it still works.

Schema 11 -> 12, and the first column this schema has REMOVED rather than added or
renamed. Removing one is easy; not breaking the files that already carry it is the part
worth testing, because the failure would not appear until somebody tried to save.

  1. THE COLUMN IS GONE
     1a. it is not in the schema, and not in a workbook this application writes
     1b. neither Monthly estimation panel offers it
     1c. nothing stamps it any more - the switch, the fill button and the gap dialog
         all used to, and a row they create now carries no such field
     1d. and the engine no longer carries manual_at on a line, which nothing read

  2. A SCHEMA 11 WORKBOOK STILL OPENS
     2a. it loads, with the same projects and people
     2b. the column is dropped rather than carried, so no row object has it
     2c. the reader is TOLD, once, and the message says why
     2d. the stated figures themselves are untouched - this is a column going, not data
     2e. every FTE is identical to the same plan without the column

  3. AND IT STILL SAVES
     3a. the retired column is not written back out, empty. rawToRows uses the header
         the FILE had, so a retired column left in that list is resurrected in a file
         stamped schema 12 - found by testing the SAVE path, not the load path
     3b. the figures go through unchanged
     3c. and an old .prap.json still imports: that reader REFUSES any column outside
         the schema, so it would have rejected the file outright

    python tools/test_retired.py
"""

import pathlib
import re
import shutil
import sys
import tempfile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parents[1]
APP = (ROOT / "app" / "PRAP.html").as_uri()
DUMMY = ROOT / "templates" / "PRAP_SourceData_Dummy_10x10_v1.11.xlsx"
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

fails = []


def check(ok, label, detail=""):
    print(f"  {'ok  ' if ok else 'FAIL'} {label}{'   ' + detail if detail else ''}")
    if not ok:
        fails.append(label)


def make_v11(dst):
    """The delivered v12 dummy, put back the way a schema 11 file looked: the column
    restored to MonthlyEstimate with a value in it, and schema_version back to 11.

    Built rather than committed, so it cannot drift from the fixture it is derived
    from - the point of the test is the DIFFERENCE, and a stale copy would stop being
    the same plan."""
    shutil.copy(DUMMY, dst)
    wb = load_workbook(dst)
    ws = wb["MonthlyEstimate"]
    hdr = [c.value for c in ws[1]]
    col = len(hdr) + 1
    ws.cell(1, col, "edited_at")
    stamped = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value:
            # exactly what the old application wrote: a JS Date, stringified
            ws.cell(r, col, "Sat Sep 12 2026 10:58:24 GMT+0000 (Coordinated Universal Time)")
            stamped += 1
    cfg = wb["Config"]
    for r in range(2, cfg.max_row + 1):
        if cfg.cell(r, 1).value == "schema_version":
            cfg.cell(r, 2, 11)
    wb.save(dst)
    return stamped


with sync_playwright() as pw:
    browser = pw.chromium.launch(executable_path=CHROME)
    pg = browser.new_page(viewport={"width": 1600, "height": 1000})
    errors = []
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.goto(APP)
    pg.wait_for_timeout(200)
    pg.set_input_files("#picker", str(DUMMY))
    pg.wait_for_timeout(4500)

    # ---------------------------------------------------------------- 1
    print("app/PRAP.html — 1. the column is gone")
    gone = pg.evaluate("""() => ({
      schema: SCHEMA_EXPECTED,
      cols: SHEET_HEADERS.MonthlyEstimate,
      retired: Object.keys(RETIRED_COLS.MonthlyEstimate || {}),
      label: typeof COLUMN_LABEL === 'object' && 'edited_at' in COLUMN_LABEL,
      help: typeof COLUMN_HELP === 'object' && 'edited_at' in COLUMN_HELP,
    })""")
    # The schema goes on stepping - 12 when edited_at was retired, 13 when milestone
    # highlighting arrived - and what this check is about is the COLUMN, not the number
    # of the day. Read the number from the file that owns it; a literal here fails on
    # the next schema change and says nothing about edited_at when it does.
    core_schema = int(re.search(r"SCHEMA_EXPECTED = (\d+);",
                                (ROOT / "src" / "core" / "00_meta.js").read_text(encoding="utf-8")
                                ).group(1))
    check(gone["schema"] == core_schema and "edited_at" not in gone["cols"],
          f"1a. not in the schema, which is at v{core_schema} and no longer names it",
          f"schema v{gone['schema']}, MonthlyEstimate: {', '.join(gone['cols'])}")
    check(gone["retired"] == ["edited_at"] and not gone["label"] and not gone["help"],
          "    and recorded as retired rather than merely deleted",
          f"RETIRED_COLS.MonthlyEstimate = {gone['retired']}")

    pg.click('nav button[data-tab="t-pers"]')
    pg.wait_for_timeout(900)
    pg.evaluate("""() => { S.selPers = 'PSN-001'; S.selAsg = 'ASG-001'; renderKeepingTab(); }""")
    pg.wait_for_timeout(900)
    panels = pg.evaluate("""() => [...document.querySelectorAll('table.data-t[data-sheet="MonthlyEstimate"]')]
        .map(t => [...t.querySelectorAll('thead th[data-cid]')].map(x => x.dataset.cid))""")
    check(panels and all("edited_at" not in p for p in panels),
          "1b. and no Monthly estimation panel offers it",
          "; ".join(", ".join(p) for p in panels))

    # A row the application creates itself carries no such field.
    made = pg.evaluate("""() => {
      const before = (S.model.raw.MonthlyEstimate || []).length;
      const a = (S.model.raw.Assignment || []).find(x => x.person_id === 'PSN-002');
      switchEstimation('assignment', a.assignment_id);
      return {aid: a.assignment_id, before};
    }""")
    pg.wait_for_timeout(400)
    try:
        pg.click("#estYes")
        pg.wait_for_timeout(1500)
    except Exception:
        pass
    stamped = pg.evaluate(f"""() => {{
      const rows = (S.model.raw.MonthlyEstimate || [])
        .filter(r => r.ref_id === '{made["aid"]}');
      return {{n: rows.length, withCol: rows.filter(r => 'edited_at' in r).length}};
    }}""")
    check(stamped["n"] > 0 and stamped["withCol"] == 0,
          "1c. and a row the switch creates carries no edited_at",
          f"{stamped['n']} month(s) seeded, {stamped['withCol']} carrying the field")
    check(pg.evaluate("() => S.calc.lines.filter(L => 'manual_at' in L).length") == 0,
          "1d. and no calculation line carries manual_at, which nothing read")

    # ---------------------------------------------------------------- 2
    print("\napp/PRAP.html — 2. a schema 11 workbook still opens")
    tmp = pathlib.Path(tempfile.mkdtemp())
    old = tmp / "PRAP_SourceData_schema11.xlsx"
    n = make_v11(old)
    print(f"       (built a schema 11 file: {n} MonthlyEstimate rows carry edited_at)")

    pg.goto(APP)
    pg.wait_for_timeout(300)
    pg.set_input_files("#picker", str(old))
    pg.wait_for_timeout(5000)
    loaded = pg.evaluate("""() => ({
      projects: Object.keys(S.model.projects).length,
      people: Object.keys(S.model.people).length,
      est: (S.model.raw.MonthlyEstimate || []).length,
      withCol: (S.model.raw.MonthlyEstimate || []).filter(r => 'edited_at' in r).length,
      told: (S.model.findings || [])
        .filter(f => String(f.msg).includes('edited_at'))
        .map(f => ({sev: f.sev, rule: f.rule, msg: f.msg})),
      total: S.calc.lines.reduce((t, L) => t + L.fte, 0),
      stated: (S.model.raw.MonthlyEstimate || []).reduce((t, r) => t + (+r.fte || 0), 0),
    })""")
    check(loaded["projects"] > 0 and loaded["people"] > 0 and loaded["est"] > 0,
          "2a. it loads, with its projects, people and stated months",
          f"{loaded['projects']} projects, {loaded['people']} people, "
          f"{loaded['est']} stated month(s)")
    check(loaded["withCol"] == 0,
          "2b. and the column is dropped, not carried into the row objects",
          f"{loaded['withCol']} row(s) still carrying it")
    told = loaded["told"]
    check(len(told) == 1 and told[0]["sev"] == "information",
          "2c. the reader is told once, as information rather than a problem",
          f"{len(told)} finding(s)" + (f" — {told[0]['rule']} {told[0]['sev']}" if told else ""))
    check(told and "change log" in told[0]["msg"],
          "    and the message says where the answer lives now",
          told[0]["msg"][:130] if told else "")

    # The same plan, opened from the delivered v12 file, must give the same figures:
    # a column going away is not allowed to move a number.
    pg.goto(APP)
    pg.wait_for_timeout(300)
    pg.set_input_files("#picker", str(DUMMY))
    pg.wait_for_timeout(5000)
    now = pg.evaluate("""() => ({
      total: S.calc.lines.reduce((t, L) => t + L.fte, 0),
      stated: (S.model.raw.MonthlyEstimate || []).reduce((t, r) => t + (+r.fte || 0), 0),
      lines: S.calc.lines.length,
    })""")
    check(abs(loaded["stated"] - now["stated"]) < 1e-9,
          "2d. the stated figures are untouched — a column went, not data",
          f"{loaded['stated']:.2f} FTE stated either way")
    check(abs(loaded["total"] - now["total"]) < 1e-9,
          "2e. and every calculated FTE is identical",
          f"{now['total']:.2f} FTE over {now['lines']} lines, both files")

    # ---------------------------------------------------------------- 3
    print("\napp/PRAP.html — 3. and it still saves")
    pg.goto(APP)
    pg.wait_for_timeout(300)
    pg.set_input_files("#picker", str(old))
    pg.wait_for_timeout(5000)
    # THE SAVE PATH, which is where this nearly went wrong. rawToRows writes the header
    # the FILE had, not the schema's - so that the application writes back a workbook
    # shaped like the one it read - and a retired column left in that list comes back
    # from the dead: written out again, empty, in a file stamped schema 12.
    out = pg.evaluate("""() => {
      const rows = rawToRows('MonthlyEstimate');
      return {header: rows[0] || [], rows: rows.length - 1,
              stated: rows.slice(1).reduce((t, r) => t + (+r[3] || 0), 0),
              headers: S.headers.MonthlyEstimate};
    }""")
    check("edited_at" not in (out.get("header") or []),
          "3a. saving an old file does not write the retired column back out, empty",
          ", ".join(out.get("header") or []))
    check("edited_at" not in (out.get("headers") or []),
          "    because it is filtered out of the headers taken from the file",
          ", ".join(out.get("headers") or []))
    check(out.get("rows", 0) > 0 and abs(out.get("stated", 0) - loaded["stated"]) < 1e-9,
          "3b. and the stated figures go through unchanged",
          f"{out.get('rows', 0)} row(s), {out.get('stated', 0):.2f} FTE")

    # And the JSON interchange reader, which REFUSES any column outside the schema -
    # so an old .prap.json would have been rejected outright rather than merely untidy.
    j = pg.evaluate("""() => {
      const doc = JSON.parse(buildPrapJson(
        Object.fromEntries(REQUIRED_SHEETS.map(s => [s, rawToRows(s)]))));
      // put the retired column back, the way a file exported by the old build has it
      for (const r of doc.sheets.MonthlyEstimate) r.edited_at = 'Sat Sep 12 2026 10:58:24 GMT+0000';
      try {
        const sheets = readPrapJson(JSON.stringify(doc));
        const me = sheets.MonthlyEstimate || [];
        return {ok: true, header: me[0] || [], rows: me.length - 1};
      } catch (e) { return {ok: false, err: String(e.message || e)}; }
    }""")
    check(j.get("ok"),
          "3c. and an old .prap.json still imports rather than being refused",
          j.get("err", "")[:150])
    check(j.get("ok") and "edited_at" not in (j.get("header") or []),
          "    with the column dropped on the way in",
          ", ".join(j.get("header") or []))

    check(not errors, "the page raised no script error", "; ".join(errors[:2]))
    shutil.rmtree(tmp, ignore_errors=True)
    browser.close()

print()
if fails:
    print("FAILURES: " + "; ".join(fails))
    sys.exit(1)
print("FAILURES: none")
