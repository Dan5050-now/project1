"""A month a manager has looked at is not the same as one nobody has (R-62).

Short, over and not staffed are findings. Some of them are decisions - "PRJ-001 is over
in the run-up to the interim lock, the sponsor is paying for it" - and drawing those in the
same alarm colour as the ones nobody has examined hides the real alarms among them. So a
month can carry a REVIEW, kept on the IssueReview sheet: 'Confirmed - no issue' and
'Accepted' close it, 'To be fixed' keeps it open with a note.

What is checked:

  * The fixture's reviews are read, and a closed month is drawn MUTED in Resource by
    project - still marked, its reason in the pop-up - while an open one is not.
  * The counts on the control, the tile and the list are OPEN issues, with the closed ones
    counted apart.
  * Clicking an issue month opens it, with the review form; a closing decision without a
    reason is refused; with one it is recorded on the whole run, with who and when.
  * A review can be changed later and removed, the same way.
  * A not-staffed month is reviewed the same way.
  * The list filters by review, and so does Resource by project's Show.
  * A REVIEW COVERS THE FIGURES IT WAS MADE ON: when the gap moves it is open again and
    V-40 says why; when the issue is gone the review is kept as history and V-40 says so.
  * The reviews travel: exported, read back by the Python reference, and judged by it
    exactly as the browser judges them.
  * A workbook from before schema 16 still opens, with no reviews and nothing refused.

    python tools/test_review.py
"""

import pathlib
import sys
import tempfile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parents[1]
APP = (ROOT / "app" / "PRAP.html").as_uri()
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
FIX = ROOT / "templates" / "PRAP_SourceData_Dummy_10x10_v1.14.xlsx"
OLDER = ROOT / "templates" / "PRAP_SourceData_Dummy_10x10_v1.13.xlsx"
TMP = pathlib.Path(tempfile.mkdtemp(prefix="prap_review_"))
P = '#t-overall .panel[data-panel="table-proj"]'

sys.path.insert(0, str(ROOT / "tools"))
import prap_io                                                         # noqa: E402

fails = []


def check(ok, label, detail=""):
    print(f"  {'ok  ' if ok else 'FAIL'} {label}{'   ' + detail if detail else ''}")
    if not ok:
        fails.append(label)


def K(y, m):
    return y * 12 + m - 1


def load(pg, path):
    pg.set_input_files("#picker", str(path))
    pg.wait_for_timeout(600)
    if pg.evaluate("() => !!(document.getElementById('replace') || {}).open"):
        pg.click("#rpYes")
    pg.wait_for_timeout(2800)


def cell(pid, k, unstaffed=False):
    return f'{P} td[data-gap="{pid}"][data-gk="{k}"]' + ('.unal' if unstaffed else '')


def classes(pg, pid, k):
    return pg.evaluate("s => (document.querySelector(s) || {}).className || ''", cell(pid, k))


def reviews(pg):
    return pg.evaluate("""() => S.model.raw.IssueReview.map(r => ({pid: r.project_id,
        mm: String(r.month).slice(0, 7), issue: r.issue, status: r.status,
        why: r.rationale, by: r.reviewed_by, gap: r.gap_fte, at: r.reviewed_at}))""")


def open_cell(pg, pid, k):
    pg.locator(cell(pid, k)).scroll_into_view_if_needed()
    pg.click(cell(pid, k))
    pg.wait_for_timeout(700)


def close_dlg(pg):
    if pg.evaluate("el('gapdlg').open"):
        pg.click("#gapClose")
        pg.wait_for_timeout(300)


def counts(pg):
    """What the calculation says is open and closed, independent of any drawing."""
    return pg.evaluate("""() => { const o = {short:0, over:0, unstaffed:0, closed:0};
        for (const r of gapRows(activeProjects())){
          if (r.rev && r.rev.closed) o.closed++; else o[r.dir]++; }
        return o; }""")


# The fixture with one project nobody is on, so a not-staffed month is there to review.
src0 = prap_io.read_xlsx(FIX)
gapped = {pid for (pid, _k) in prap_io.calculate(prap_io.Model(prap_io.read_xlsx(FIX)))["proj_gap"]}
GONE = next(p["project_id"] for p in src0["Project"]
            if p["project_id"] not in gapped and str(p.get("status") or "") != "Completed"
            and any(a["project_id"] == p["project_id"] for a in src0["Assignment"]))
src0["Assignment"] = [a for a in src0["Assignment"] if a["project_id"] != GONE]
keep = {a["assignment_id"] for a in src0["Assignment"]}
for sh in ("PersonPeriodWeight", "MonthlyEstimate"):
    src0[sh] = [r for r in src0.get(sh, []) if not r.get("assignment_id") or r["assignment_id"] in keep]
PLAN = TMP / "plan.xlsx"
prap_io.write_xlsx(src0, PLAN)

OCT, NOV, DEC, JAN, FEB = K(2026, 10), K(2026, 11), K(2026, 12), K(2027, 1), K(2027, 2)

print("R-62 — a manager's review of a month off its standard")
with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path=CHROME, downloads_path=str(TMP))
    pg = b.new_context(accept_downloads=True, viewport={"width": 1560, "height": 1000}).new_page()
    pg.set_default_timeout(25000)
    errors = []
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.on("dialog", lambda d: d.accept())
    pg.goto(APP)
    pg.wait_for_timeout(400)
    load(pg, PLAN)

    print("\n1. the fixture's reviews, and what they look like")
    rv = reviews(pg)
    check(len(rv) == 3 and {(r["mm"], r["issue"], r["status"]) for r in rv}
          == {("2026-10", "over", "Accepted"), ("2026-12", "over", "Accepted"),
              ("2026-11", "short", "To be fixed")},
          "the three example reviews are read", str([(r["mm"], r["status"]) for r in rv]))
    check(pg.evaluate("S.model.findings.filter(f => f.rule === 'V-40').length") == 0,
          "and each still describes its month, so V-40 says nothing")
    muted = pg.eval_on_selector_all(f"{P} td.c.rev", "es => es.map(e => +e.dataset.gk)")
    check(sorted(muted) == [OCT, DEC],
          "THE TWO ACCEPTED MONTHS ARE DRAWN MUTED, and only those", str(muted))
    check("gapc over" in classes(pg, "PRJ-001", OCT),
          "still marked as over - muted is not unmarked", classes(pg, "PRJ-001", OCT))
    nov = classes(pg, "PRJ-001", NOV)
    check("gapc short" in nov and " rev" not in nov,
          "'To be fixed' keeps the month open, in its alarm colour", nov)
    tip = pg.get_attribute(cell("PRJ-001", OCT), "data-tip") or ""
    check("Sponsor-funded" in tip and "Accepted" in tip,
          "the pop-up carries the decision and its reason", tip[-140:])

    print("\n2. the counts are open issues, the closed ones counted apart")
    c = counts(pg)
    btn = pg.inner_text(f"{P} .gapbtn").replace("\n", " ")
    check(f"{c['short']} short" in btn and f"{c['over']} over" in btn
          and f"{c['closed']} reviewed" in btn and c["closed"] == 2,
          "THE CONTROL COUNTS WHAT IS STILL OPEN, and says how many were reviewed", btn)
    tile = pg.evaluate("""() => [...document.querySelectorAll('#t-overall .tile')]
        .find(t => t.textContent.includes('Off their standard')).textContent""")
    check(f"{c['closed']} reviewed" in tile, "and so does the tile", " ".join(tile.split())[:120])

    print("\n3. clicking a month opens it, with the review")
    open_cell(pg, "PRJ-001", OCT)
    box = pg.inner_text("#gapdlg .revbox") if pg.evaluate("el('gapdlg').open") else ""
    check("Accepted" in box and "Sponsor-funded" in box,
          "CLICKING A REVIEWED MONTH SHOWS THE DECISION AND ITS REASON", box[:110].replace("\n", " "))
    check(pg.evaluate("el('gapdlg').querySelector('.revbox').classList.contains('closed')"),
          "drawn as closed")
    close_dlg(pg)

    open_cell(pg, "PRJ-001", JAN)
    check("Not reviewed yet" in pg.inner_text("#gapdlg .revbox"),
          "an open month says it has not been reviewed")
    before, pend = len(reviews(pg)), pg.evaluate("S.pending.length")
    pg.select_option("#revStatus", "Accepted")
    pg.fill("#revWhy", "")
    pg.click("#gapdlg [data-revsave]")
    pg.wait_for_timeout(500)
    check(len(reviews(pg)) == before and pg.evaluate("S.pending.length") == pend
          and pg.evaluate("el('gapdlg').open"),
          "A CLOSING DECISION WITHOUT A REASON IS REFUSED - nothing recorded, the form stays")

    run = pg.evaluate(f"issueRun('PRJ-001', {JAN}, 'over')")
    gaps = pg.evaluate(f"""() => Object.fromEntries({run}.map(k => [k,
        Math.round(gapOf('PRJ-001', k).gap * 100) / 100]))""")
    pg.fill("#revWhy", "Planned surge before the interim lock.")
    pg.fill("#revBy", "Lee M.")
    check(pg.is_checked("#revRun"), "the whole run is offered, and ticked", str(run))
    pg.click("#gapdlg [data-revsave]")
    pg.wait_for_timeout(900)
    rv = {r["mm"]: r for r in reviews(pg) if r["issue"] == "over"}
    mms = [f"{k // 12}-{k % 12 + 1:02d}" for k in run]
    check(len(run) >= 2 and all(rv.get(m, {}).get("status") == "Accepted"
                                and rv[m]["why"] == "Planned surge before the interim lock."
                                and rv[m]["by"] == "Lee M." and rv[m]["at"] for m in mms),
          "RECORDED ON EVERY MONTH OF THE RUN, with who and when", ", ".join(mms))
    check(all(abs(rv[m]["gap"] - gaps[str(k)]) < 1e-9 for m, k in zip(mms, run)),
          "each with the gap it was reviewed at", str(gaps))
    check(len(reviews(pg)) == before + len([m for m in mms if m != "2026-12"]),
          "a month already reviewed is updated, not given a second review")
    check(sorted(pg.eval_on_selector_all(f"{P} td.c.rev", "es => es.map(e => +e.dataset.gk)"))
          == sorted(set([OCT] + run)), "and every one of them is muted now")
    check(any(p["sheet"] == "IssueReview" for p in pg.evaluate("S.pending")),
          "it is an edit like any other, waiting for Save")
    box = pg.inner_text("#gapdlg .revbox")
    check("Accepted" in box and "Lee M." in box and "Planned surge" in box,
          "and the open dialog shows it at once", box[:90].replace("\n", " "))
    c2 = counts(pg)
    check(c2["over"] == c["over"] - (len(run) - 1) and c2["closed"] == c["closed"] + len(run) - 1,
          "the open count falls by what was closed", f"{c['over']} -> {c2['over']} over")

    print("\n4. changed later, and removed")
    pg.select_option("#revStatus", "To be fixed")
    pg.fill("#revWhy", "Move one CRA off in Q1.")
    pg.uncheck("#revRun")
    pg.click("#gapdlg [data-revsave]")
    pg.wait_for_timeout(800)
    jan = classes(pg, "PRJ-001", JAN)
    check("gapc over" in jan and " rev" not in jan,
          "CHANGED TO 'TO BE FIXED', THE MONTH IS OPEN AGAIN - and only that month", jan)
    check(" rev" in classes(pg, "PRJ-001", FEB), "its neighbours keep their own decision")
    check("To be fixed" in pg.inner_text("#gapdlg .revbox"), "the dialog says so")
    close_dlg(pg)
    open_cell(pg, "PRJ-001", FEB)
    check(not pg.is_checked("#revRun"),
          "a month already reviewed offers its run UNTICKED - a change touches this month "
          "unless asked, since its neighbours may hold decisions of their own")
    pg.check("#revRun")
    pg.click("#gapdlg [data-revclear]")
    pg.wait_for_timeout(900)
    left = sorted((r["mm"], r["issue"]) for r in reviews(pg))
    check(left == [("2026-10", "over"), ("2026-11", "short")],
          "REMOVE TAKES THE REVIEW OFF THE WHOLE RUN, as Save put it on", str(left))
    check(" rev" not in classes(pg, "PRJ-001", FEB), "and the month is open again")
    close_dlg(pg)

    print("\n5. a not-staffed month is reviewed the same way")
    un = pg.evaluate(f"""() => {{ const q = [...document.querySelectorAll('{P} td.c.unal')][0];
        return q ? [q.dataset.gap, +q.dataset.gk] : null; }}""")
    check(un is not None, "there is a not-staffed month on screen", str(un))
    if un:
        n_un = counts(pg)["unstaffed"]
        pg.locator(cell(un[0], un[1], True)).scroll_into_view_if_needed()
        pg.click(cell(un[0], un[1], True))
        pg.wait_for_timeout(700)
        check(pg.evaluate("el('gapdlg').open") and pg.locator("#gapdlg .revbox").count() == 1,
              "clicking it opens the month with its review", pg.inner_text("#gapTitle"))
        pg.select_option("#revStatus", "Confirmed - no issue")
        pg.fill("#revWhy", "Start-up handled by the sponsor's own team.")
        if pg.locator("#revRun").count():
            pg.uncheck("#revRun")
        pg.click("#gapdlg [data-revsave]")
        pg.wait_for_timeout(800)
        close_dlg(pg)
        cl = pg.evaluate("s => document.querySelector(s).className", cell(un[0], un[1], True))
        check(" rev" in cl and counts(pg)["unstaffed"] == n_un - 1,
              "CONFIRMED, IT IS MUTED AND LEAVES THE OPEN COUNT", cl)

    print("\n6. the list, and Show, filter by review")
    pg.click(f"{P} .gapbtn")
    pg.wait_for_timeout(700)
    pg.click('#gapsdlg [data-gaprev="closed"]')
    pg.wait_for_timeout(500)
    rows = pg.eval_on_selector_all("#gapsdlg tr.gaprow", "es => es.map(e => e.className)")
    closed_n = counts(pg)["closed"]
    check(rows and len(rows) == min(closed_n, 40) and all(" rev" in r for r in rows),
          "REVIEWED shows the closed months, and only them", f"{len(rows)} of {closed_n}")
    check("Sponsor-funded" in pg.inner_text("#gapsBody"), "with their reasons in the list")
    pg.click('#gapsdlg [data-gaprev="open"]')
    pg.wait_for_timeout(500)
    rows = pg.eval_on_selector_all("#gapsdlg tr.gaprow", "es => es.map(e => e.className)")
    check(rows and not any(" rev" in r for r in rows), "OPEN leaves them out")
    pg.click("[data-gapclear]")
    pg.wait_for_timeout(300)
    pg.click("#gapsClose")
    pg.wait_for_timeout(300)
    pg.select_option("#projIssue", "reviewed")
    pg.wait_for_timeout(600)
    shown = pg.eval_on_selector_all(f"{P} tr.parent", "es => es.map(e => e.dataset.k.slice(2))")
    want = sorted({r["pid"] for r in reviews(pg) if r["status"] != "To be fixed"})
    check(sorted(shown) == want,
          "Show 'Reviewed and closed' lists the projects with a closed month", str(shown))
    pg.select_option("#projIssue", "")
    pg.wait_for_timeout(500)

    print("\n7. a review covers the figures it was made on (V-40)")
    msgs = pg.evaluate("""() => {
        const r = S.model.raw.IssueReview.find(x => String(x.month).startsWith('2026-10'));
        r.gap_fte = 0.5;
        S.model.raw.IssueReview.push({__row: 9001, project_id: 'PRJ-001', month: '2026-11',
          issue: 'over', status: 'Accepted', gap_fte: 0.3, rationale: 'x',
          reviewed_by: null, reviewed_at: null});
        rebuild(true); renderKeepingTab();
        return S.model.findings.filter(f => f.rule === 'V-40').map(f => f.msg); }""")
    pg.wait_for_timeout(600)
    check(any("2026-10" in m and "fresh look" in m for m in msgs),
          "A GAP THAT HAS MOVED SINCE THE REVIEW IS SAID TO NEED A FRESH LOOK", str(msgs)[:150])
    check(any("2026-11" in m and "no longer over" in m for m in msgs),
          "and a review of an issue that has gone is kept as history, and said so")
    octc = classes(pg, "PRJ-001", OCT)
    check("gapc over" in octc and " rev" not in octc,
          "THE MONTH IS OPEN AGAIN - a decision about other numbers does not close it", octc)
    open_cell(pg, "PRJ-001", OCT)
    check("Re-check" in pg.inner_text("#gapdlg .revbox"), "and its dialog asks for a re-check")
    close_dlg(pg)

    print("\n8. the reviews travel")
    pg.evaluate("() => { const r = S.model.raw.IssueReview.find(x => x.__row === 9001); "
                "S.model.raw.IssueReview.splice(S.model.raw.IssueReview.indexOf(r), 1); "
                "rebuild(true); renderKeepingTab(); }")
    js_v40 = sorted(pg.evaluate("S.model.findings.filter(f => f.rule === 'V-40').map(f => f.msg)"))
    pg.click("#saveBtn")                 # a review is an edit: kept by Save, like any other
    pg.wait_for_timeout(1600)
    for sel in ("#cfYes", "#saveAnyway"):
        if pg.locator(sel).count() and pg.locator(sel).is_visible():
            pg.click(sel)
            pg.wait_for_timeout(1200)
    check(pg.evaluate("S.pending.length") == 0
          and any(a.get("sheet") == "IssueReview" for a in pg.evaluate("S.audit")),
          "SAVE KEEPS THE REVIEWS, and the change log records them")
    with pg.expect_download() as dl:
        pg.click("#exportBtn")
        pg.click("#exportBtn2")
    out = TMP / "reviewed.xlsx"
    dl.value.save_as(out)
    wb = load_workbook(out, read_only=True)
    hdr = [c.value for c in next(wb["IssueReview"].iter_rows(max_row=1))] \
        if "IssueReview" in wb.sheetnames else []
    check(hdr[:8] == ["project_id", "month", "issue", "status", "gap_fte", "rationale",
                      "reviewed_by", "reviewed_at"],
          "THE EXPORT HAS THE ISSUEREVIEW SHEET, under its own headings", str(hdr[:4]))
    src = prap_io.read_xlsx(out)
    got = sorted((r["project_id"], str(r["month"])[:7], r["issue"], r["status"])
                 for r in src.get("IssueReview", []))
    app = sorted((r["pid"], r["mm"], r["issue"], r["status"]) for r in reviews(pg))
    check(got == app, "and the Python reference reads back exactly what the browser holds",
          f"{len(got)} review(s)")
    M = prap_io.Model(src)
    prap_io.calculate(M)
    py_v40 = sorted(f["msg"] for f in M.findings if f.get("rule") == "V-40")
    check(py_v40 == js_v40 and len(js_v40) >= 1,
          "BOTH ENGINES JUDGE THE REVIEWS THE SAME, message for message", f"{len(py_v40)} V-40")

    print("\n9. a workbook from before schema 16")
    if OLDER.exists():
        load(pg, OLDER)
        st = pg.evaluate("""() => ({rows: S.model.raw.IssueReview.length,
            fatal: S.model.findings.filter(f => f.sev === 'fatal').length,
            v09: S.model.findings.some(f => f.rule === 'V-09' && f.sheet === 'IssueReview'),
            rev: document.querySelectorAll('td.c.rev').length})""")
        check(st["rows"] == 0 and st["fatal"] == 0 and st["v09"] and st["rev"] == 0,
              "STILL OPENS - no reviews, said once as information, nothing refused", str(st))
        open_cell(pg, *pg.evaluate(f"""() => {{ const q = document.querySelector('{P} td[data-gap]');
            return [q.dataset.gap, +q.dataset.gk]; }}"""))
        check(pg.locator("#gapdlg .revbox").count() == 1,
              "and its issues can be reviewed from the first click")
        close_dlg(pg)

    check(not errors, "no script error anywhere in the run", "; ".join(errors[:3]))
    b.close()

print(f"\nFAILURES: {'none' if not fails else len(fails)}")
for f in fails:
    print(f"  FAILED  {f}")
sys.exit(1 if fails else 0)
